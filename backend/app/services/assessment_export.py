"""운영평가 회차 엑셀 내보내기 (2026-10-13, 13.9-122) — **통제 1건 = 시트 1장**, 증빙 이미지를 시트에 붙인다.

회사가 쓰던 운영평가 파일(HR 운영평가 20251130.xlsx)의 모양을 따른다:
- 1~3행 머리 — 샘플수·프로세스·하위프로세스·위험(번호·내용)·담당자·통제활동이름·통제현황
- "증빙자료" — 증빙마다 문서명 | 비고 한 줄, 그 아래 이미지를 세로로 차례대로
- 이미지(PNG·JPG)는 그대로, **PDF 는 페이지마다 이미지로**(pypdfium2 — Apache/BSD. PyMuPDF 는 AGPL 이라 쓰지 않는다,
  상용화 §20). PDF 는 1개당 최대 10페이지(넘으면 "이후 n페이지는 원본 참조"). 그 밖의 파일은 파일명과 안내만.

Regina 결정(STEP 0): 회차 전체 또는 내 담당(통제책임자)만 / 샘플수 = 통제 주기 권장 표본 수 / 통제현황(WTT) = 통제 설명,
기존 프로세스 칸은 비움 / 문서명·비고 = 증빙의 title·note(없으면 파일명·올린 사람과 일시) / 증빙 없는 통제도 시트를 만든다.
"""
from __future__ import annotations

import io
import math
from datetime import datetime
from uuid import UUID

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.models.assessment import (
    ACTIVITY_OPERATION_ASSESSMENT,
    ACTIVITY_TEST,
    AssessmentActivity,
    AssessmentCycle,
    CycleTarget,
)
from app.models.evidence import EvidenceFile
from app.models.role_assignment import ROLE_CONTROL_OWNER
from app.models.user import User
from app.services import control_resolver as cr
from app.services import sampling
from app.services.role_resolver import _assignments_by_target, resolve_roles_for_control

HEADERS = ("샘플수", "프로세스번호", "프로세스이름", "하위프로세스번호", "하위프로세스이름", "위험번호", "담당자",
           "위험내용", "통제활동이름", "통제현황(WTT의 summary 수준)", "통제현황(기존프로세스)")
WIDTHS = (8, 12, 16, 16, 18, 14, 10, 40, 24, 50, 30)
RESULT_LABELS = {"pass": "적정", "fail": "미흡", "na": "해당 없음"}
MAX_IMAGE_WIDTH = 760       # px — A~J 열 폭 안에 들어오게
ROW_PX = 20                 # 기본 행 높이(15pt)
PDF_MAX_PAGES = 10
PDF_SCALE = 150 / 72        # 약 150dpi
IMAGE_MIME = {"image/png", "image/jpeg"}

_BOLD = Font(bold=True)
_HEAD_FILL = PatternFill("solid", fgColor="DDE3EE")
_SECTION_FILL = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="BBBBBB")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")


class ExportError(Exception):
    pass


def _sheet_title(code: str, used: set[str]) -> str:
    base = "".join(ch for ch in (code or "통제") if ch not in '[]:*?/\\')[:31] or "통제"
    t, i = base, 2
    while t in used:
        t = f"{base[:28]}_{i}"
        i += 1
    used.add(t)
    return t


def _fetch(key: str) -> bytes:
    from app.minio_client import get_object_stream
    resp = get_object_stream(key)
    try:
        return resp.read()
    finally:
        resp.close()
        resp.release_conn()


def _pdf_pages(data: bytes) -> tuple[list[bytes], int]:
    """PDF → 페이지 PNG(최대 10장), 전체 페이지 수."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(data)
    try:
        total = len(pdf)
        out = []
        for i in range(min(total, PDF_MAX_PAGES)):
            page = pdf[i]
            pil = page.render(scale=PDF_SCALE).to_pil()
            buf = io.BytesIO()
            pil.save(buf, format="PNG")
            out.append(buf.getvalue())
            page.close()
        return out, total
    finally:
        pdf.close()


def _place_image(ws, data: bytes, row: int) -> int:
    """이미지를 row 행 A열에 붙이고, 다음에 쓸 행 번호를 돌려준다."""
    from PIL import Image as PILImage
    with PILImage.open(io.BytesIO(data)) as im:
        w, h = im.size
    img = XLImage(io.BytesIO(data))
    if w > MAX_IMAGE_WIDTH:
        img.width, img.height = MAX_IMAGE_WIDTH, int(h * MAX_IMAGE_WIDTH / w)
    else:
        img.width, img.height = w, h
    ws.add_image(img, f"A{row}")
    return row + math.ceil(img.height / ROW_PX) + 1


def _fmt_time(t: datetime | None) -> str:
    return t.astimezone().strftime("%Y-%m-%d %H:%M") if t else ""


def build(db: Session, cycle: AssessmentCycle, *, owner_id: UUID | None = None,
          fetch=_fetch) -> tuple[bytes, int]:
    """회차 → xlsx 바이트, 시트 수. `owner_id` 가 있으면 그 사람이 통제책임자인 통제만. `fetch` 는 테스트용 주입점."""
    targets = {t.control_id for t in db.query(CycleTarget).filter(CycleTarget.cycle_id == cycle.id).all()}
    processes, sub_processes, risks = cr.resolve_hierarchy(db)
    proc = {p["code"]: p for p in processes}
    subp = {s["code"]: s for s in sub_processes}
    risk = {r["id"]: r for r in risks}
    controls = sorted((c for c in cr.resolve_controls(db) if c["id"] in targets), key=lambda c: c["code"])
    if owner_id is not None:
        assignments = _assignments_by_target(db)
        proc_ids = {p["code"]: p["id"] for p in processes}
        controls = [c for c in controls if any(
            r["role_name"] == ROLE_CONTROL_OWNER and r["user_id"] == owner_id
            for r in resolve_roles_for_control(db, c["id"], proc_ids.get(c.get("process_code")), assignments=assignments))]
    if not controls:
        raise ExportError("내보낼 통제가 없습니다")

    ids = [c["id"] for c in controls]
    files = db.query(EvidenceFile).filter(EvidenceFile.cycle_id == cycle.id, EvidenceFile.control_id.in_(ids),
                                          EvidenceFile.is_deleted == False).order_by(  # noqa: E712
        EvidenceFile.created_at).all()
    acts = db.query(AssessmentActivity).filter(AssessmentActivity.cycle_id == cycle.id,
                                               AssessmentActivity.control_id.in_(ids),
                                               AssessmentActivity.is_deleted == False).order_by(  # noqa: E712
        AssessmentActivity.performed_at).all()
    names = {u.id: u.display_name for u in db.query(User).filter(
        User.id.in_({f.uploaded_by_id for f in files} | {a.performed_by_id for a in acts})).all()} if (files or acts) else {}

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for c in controls:
        ws = wb.create_sheet(_sheet_title(c["code"], used))
        for i, w in enumerate(WIDTHS, start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        sp = subp.get(c.get("sub_process_code")) or {}
        p = proc.get(c.get("process_code")) or {}
        rk = risk.get(c.get("risk_id")) or {}
        values = (sampling.required_samples(cycle.kind, c), c.get("process_code"), p.get("name"),
                  c.get("sub_process_code"), sp.get("name"), rk.get("code"), c.get("owner_name"),
                  rk.get("description"), c.get("name"), c.get("description"), None)
        for col, (h, v) in enumerate(zip(HEADERS, values, strict=True), start=1):
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
            hc = ws.cell(1, col, h)
            hc.font, hc.fill, hc.alignment, hc.border = _BOLD, _HEAD_FILL, Alignment(wrap_text=True, vertical="center",
                                                                                     horizontal="center"), _BOX
            vc = ws.cell(3, col, v)
            vc.alignment, vc.border = _WRAP, _BOX
        ws.row_dimensions[3].height = 120
        ws.cell(4, 1, f"통제코드 {c['code']} · 회차 {cycle.name} · 평가 주기 {c.get('assessment_frequency') or ''}").font = \
            Font(color="666666", size=9)

        # 평가 결과 — 활동 종류별 마지막 기록
        last = {}
        for a in acts:
            if a.control_id == c["id"]:
                last[a.activity_kind] = a
        row = 5
        for kind, label in ((ACTIVITY_TEST, "자체점검"), (ACTIVITY_OPERATION_ASSESSMENT, "운영평가")):
            a = last.get(kind)
            text = (f"{RESULT_LABELS.get(a.result or '', a.result or '결과 미기록')}"
                    f"{' — ' + a.note if a.note else ''} ({names.get(a.performed_by_id, '')}, {_fmt_time(a.performed_at)})"
                    if a else "기록 없음")
            ws.cell(row, 1, label).font = _BOLD
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=11)
            ws.cell(row, 2, text).alignment = _WRAP
            row += 1

        row += 1
        sec = ws.cell(row, 1, "증빙자료")
        sec.font, sec.fill = _BOLD, _SECTION_FILL
        row += 1
        ws.cell(row, 1, "문서명").font = _BOLD
        ws.cell(row, 4, "비고").font = _BOLD
        row += 1
        mine = [f for f in files if f.control_id == c["id"]]
        if not mine:
            ws.cell(row, 1, "증빙 없음").font = Font(color="C00000", bold=True)
            continue
        for f in mine:
            ws.cell(row, 1, f.title or f.filename).font = _BOLD
            note = f.note or f"{names.get(f.uploaded_by_id, '')} · {_fmt_time(f.created_at)}"
            ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=11)
            ws.cell(row, 4, note).alignment = _WRAP
            row += 1
            if not f.minio_key:
                ws.cell(row, 1, "파일 본체 없음(레거시)")
                row += 2
                continue
            try:
                data = fetch(f.minio_key)
                if f.mime_type in IMAGE_MIME:
                    row = _place_image(ws, data, row)
                elif f.mime_type == "application/pdf":
                    pages, total = _pdf_pages(data)
                    for pg in pages:
                        row = _place_image(ws, pg, row)
                    if total > len(pages):
                        ws.cell(row, 1, f"이후 {total - len(pages)}페이지는 원본 참조 ({f.filename})").font = \
                            Font(color="666666", italic=True)
                        row += 2
                else:
                    ws.cell(row, 1, f"원본 파일({f.filename})은 시스템에서 다운로드").font = Font(color="666666", italic=True)
                    row += 2
            except Exception as e:  # noqa: BLE001 — 한 파일이 깨져도 나머지는 내보낸다
                ws.cell(row, 1, f"이미지로 붙이지 못함({f.filename}): {e}").font = Font(color="C00000")
                row += 2
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(controls)
