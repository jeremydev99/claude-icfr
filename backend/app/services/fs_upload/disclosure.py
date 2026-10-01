"""공시양식형 판독기 — DART 표준 양식 (8-B, ADR-0037 §3).

실측(2025 감사 완료본 `BS공시`·`PL공시`·`CF공시`):
- 제목 행("재 무 상 태 표"), 기간 행("제26기 2025년 12월 31일 현재"), 단위 행("(단위 : 원)")
- 머리 행: `과목` | `주석` | `제 26 기말`(2열 병합) | `제 25 기말`(2열 병합)
- 기간당 금액 열이 2개(안쪽=세부, 바깥=소계)지만 **소계가 안쪽 열에 있기도 하다**(CF) —
  열 위치로 소계를 판정하지 않는다. 위치는 `raw_meta.col` 로 보존만 한다.
- 오른쪽 검산 열(TRUE)·메모 열은 기간 범위 밖이라 읽지 않는다.

병합 정보에 기대지 않는다 — 기간 범위는 "다음 기간 머리 직전까지", 마지막 기간은 첫 기간과 같은 폭.
"""
import re

from openpyxl.worksheet.worksheet import Worksheet

from app.services.fs_upload.cells import (
    AmountError,
    detect_unit,
    display_label,
    leading_ws,
    norm,
    parse_amount,
    raw_text,
    split_prefix,
)
from app.services.fs_upload.parsed import KIND_DISCLOSURE, ParsedRow, ParsedSheet

HEAD_SCAN_ROWS = 12
# 제목 → 재무제표 종류. "포괄손익계산서"가 "손익계산서"보다 먼저 와야 하지는 않다(둘 다 PL)
TITLE_TYPES = (("재무상태표", "BS"), ("대차대조표", "BS"), ("손익계산서", "PL"),
               ("현금흐름표", "CF"), ("자본변동표", "SCE"))
_TERM_RE = re.compile(r"제(\d+)기")
_TERM_YEAR_RE = re.compile(r"제(\d+)기.*?(\d{4})년")
_YEAR_RE = re.compile(r"(\d{4})년")


def _head_cells(ws: Worksheet):
    for r in range(1, min(HEAD_SCAN_ROWS, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                yield r, c, v


def find_header(ws: Worksheet) -> tuple[int, int] | None:
    """`과목` 머리 셀 (행, 열). 같은 행에 `제 N 기` 셀이 있어야 공시양식으로 본다."""
    for r, c, v in _head_cells(ws):
        if norm(v) == "과목":
            if any(isinstance(ws.cell(r, j).value, str) and _TERM_RE.search(norm(ws.cell(r, j).value))
                   for j in range(c + 1, ws.max_column + 1)):
                return r, c
    return None


def title_type(ws: Worksheet) -> str | None:
    """표 머리 영역의 제목으로 본 재무제표 종류. 양식 판독과 별개로 미지원 양식(SCE)을 알리는 데 쓴다."""
    for _r, _c, v in _head_cells(ws):
        n = norm(v)
        for word, code in TITLE_TYPES:
            if word in n:
                return code
    return None


def looks_like(ws: Worksheet) -> bool:
    return find_header(ws) is not None


def read(ws: Worksheet) -> ParsedSheet:
    header = find_header(ws)
    if header is None:
        raise ValueError("공시양식 머리 행(과목·제 N 기)을 찾지 못했습니다")
    hdr_row, label_col = header

    stype, unit_word, unit, consolidated = None, None, None, False
    term_years: dict[int, int] = {}
    for r, _c, v in _head_cells(ws):
        n = norm(v)
        if r < hdr_row and stype is None:
            for word, code in TITLE_TYPES:
                if word in n:
                    stype = code
                    consolidated = "연결" in n
                    break
        u = detect_unit(v)
        if u and unit_word is None:
            unit_word, unit = u
        m = _TERM_YEAR_RE.search(n)
        if m:
            term_years.setdefault(int(m.group(1)), int(m.group(2)))

    # 기간 열
    notes_col = None
    starts: list[tuple[int, str]] = []
    for c in range(label_col + 1, ws.max_column + 1):
        v = ws.cell(hdr_row, c).value
        if not isinstance(v, str):
            continue
        n = norm(v)
        if n == "주석":
            notes_col = c
        elif _TERM_RE.search(n):
            starts.append((c, n))
    sheet = ParsedSheet(kind=KIND_DISCLOSURE, sheet_name=ws.title, statement_type=stype, periods=[],
                        rows=[], unit_word=unit_word, unit=unit, consolidated_hint=consolidated)
    spans: list[tuple[int, list[int]]] = []
    width = None
    for i, (c, n) in enumerate(starts):
        end = starts[i + 1][0] - 1 if i + 1 < len(starts) else c + (width or 1) - 1
        cols = list(range(c, end + 1))
        width = width or len(cols)
        year = None
        m = _YEAR_RE.search(n)
        if m:
            year = int(m.group(1))
        else:
            term = int(_TERM_RE.search(n).group(1))
            year = term_years.get(term)
        if year is None:
            sheet.errors.append(f"'{ws.cell(hdr_row, c).value}' 기간의 회계연도를 찾지 못했습니다")
            continue
        if any(y == year for y, _ in spans):
            sheet.errors.append(f"회계연도 {year} 기간 열이 두 번 나옵니다")
            continue
        spans.append((year, cols))
    sheet.periods = [y for y, _ in sorted(spans, key=lambda x: -x[0])]

    for r in range(hdr_row + 1, ws.max_row + 1):
        cell = ws.cell(r, label_col)
        v = cell.value
        if v is None or (isinstance(v, str) and not v.strip("  　\t")):
            continue
        if not isinstance(v, str) or norm(v).startswith("["):
            break
        rank, _ = split_prefix(v)
        indent = int(cell.alignment.indent or 0) * 2 + leading_ws(v)
        row = ParsedRow(row_no=r, raw_label=v, label=display_label(v), rank=rank, indent=indent,
                        bold=bool(cell.font and cell.font.b), amounts={})
        if notes_col is not None and ws.cell(r, notes_col).value is not None:
            row.meta["notes"] = str(ws.cell(r, notes_col).value)
        cols_used: dict[str, str] = {}
        for year, cols in spans:
            found: list[tuple] = []   # (열, 원본 값, 금액, 메모)
            for c in cols:
                cv = ws.cell(r, c).value
                try:
                    d, memo = parse_amount(cv)
                except AmountError as e:
                    row.errors.append(f"{year}: {e}")
                    continue
                if d is not None:
                    found.append((c, cv, d, memo))
            if len({f[2] for f in found}) > 1:
                row.errors.append(f"{year}: 한 기간에 금액이 둘입니다 ({', '.join(str(f[1]) for f in found)})")
                row.amounts[year] = None
                continue
            if found:
                c, cv, d, memo = found[-1]
                row.amounts[year] = d
                row.raw_values[year] = raw_text(cv)
                cols_used[str(year)] = "inner" if c != cols[-1] else "outer"
                if memo == "float_noise":
                    row.meta.setdefault("float_noise", []).append(year)
            else:
                row.amounts[year] = None
        if cols_used:
            row.meta["col"] = cols_used
        sheet.rows.append(row)
    return sheet
