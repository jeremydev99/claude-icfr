"""가로 연도형 판독기 — 한 시트에 여러 연도 열 (8-B, ADR-0037 §3).

실측(2025 감사 완료본 `BS정산표`·`PL정산표`):
- 1행: `과목` | 2020 | 2021 | … | 2025 (연도 정수), 2행: `Final` | … | `조정전`·`IFRS조정`·`감사수정`·`Final`
- 연도 아래 열이 여럿이면 **`Final` 열만 금액**, 나머지는 조정 내역으로 `raw_meta.adjustments` 에 보존
- 연도 범위 밖 텍스트 열(COA→공시 계정 매핑 등)은 `raw_meta.annotations` 에 보존
- 계층은 번호 접두 + 굵게(들여쓰기 없음). 단위 표기 없음 → 사용자 선택 필수
- 하단 `[영업외손익분류]` 같은 분석표·행번호 행에서 멈춘다
- **원본 결함 대응**: 연도 머리 중복(`PL정산표` 2022 두 번) → 그 연도 제외 + 경고
"""
import re

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.fs_upload.cells import (
    AmountError,
    detect_unit,
    display_label,
    leading_ws,
    norm,
    parse_amount,
    raw_text,
    split_prefix,
)
from app.services.fs_upload.parsed import KIND_HORIZONTAL, ParsedRow, ParsedSheet

HEAD_SCAN_ROWS = 10
YEAR_MIN, YEAR_MAX = 1990, 2100
FINAL_TAGS = {"final", "최종", "확정"}
_YEAR_TEXT_RE = re.compile(r"^(?:FY)?(\d{4})(?:년)?$", re.I)


def _year_of(v) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        y = v
    elif isinstance(v, float) and v.is_integer():
        y = int(v)
    elif isinstance(v, str):
        m = _YEAR_TEXT_RE.match(norm(v))
        if not m:
            return None
        y = int(m.group(1))
    else:
        return None
    return y if YEAR_MIN <= y <= YEAR_MAX else None


def find_header(ws: Worksheet) -> tuple[int, list[tuple[int, int]]] | None:
    """연도 머리 행과 (열, 연도) 목록. 연도 셀이 2개 이상인 첫 행."""
    for r in range(1, min(HEAD_SCAN_ROWS, ws.max_row) + 1):
        ys = [(c, y) for c in range(1, ws.max_column + 1) if (y := _year_of(ws.cell(r, c).value)) is not None]
        if len(ys) >= 2:
            return r, ys
    return None


def looks_like(ws: Worksheet) -> bool:
    return find_header(ws) is not None


def _label_col(ws: Worksheet, hdr_row: int, first_year_col: int) -> int:
    for c in range(1, first_year_col):
        v = ws.cell(hdr_row, c).value
        if isinstance(v, str) and norm(v) in {"과목", "계정과목", "계정"}:
            return c
    return 1


def read(ws: Worksheet) -> ParsedSheet:
    found = find_header(ws)
    if found is None:
        raise ValueError("연도 머리 행을 찾지 못했습니다")
    hdr_row, year_cells = found
    label_col = _label_col(ws, hdr_row, year_cells[0][0])

    # 상태 행(Final/조정전…) — 머리 다음 행에 문자열이 있으면
    status = {c: str(ws.cell(hdr_row + 1, c).value).strip() for c in range(1, ws.max_column + 1)
              if c != label_col and isinstance(ws.cell(hdr_row + 1, c).value, str)}
    data_start = hdr_row + (2 if status else 1)

    sheet = ParsedSheet(kind=KIND_HORIZONTAL, sheet_name=ws.title, statement_type=None, periods=[], rows=[])
    for r in range(1, hdr_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and (u := detect_unit(v)):
                sheet.unit_word, sheet.unit = u
                break

    # 연도별 열 범위 — 다음 연도 머리 직전까지, 마지막은 상태 표기가 있는 마지막 열까지
    spans: dict[int, list[list[int]]] = {}
    for i, (c, y) in enumerate(year_cells):
        if i + 1 < len(year_cells):
            end = year_cells[i + 1][0] - 1
        else:
            end = max([c] + [k for k in status if k > c and _contiguous(status, c, k)])
        spans.setdefault(y, []).append(list(range(c, end + 1)))
    periods: list[tuple[int, int, list[int]]] = []
    used_cols: set[int] = set()
    for y, lst in spans.items():
        for cols in lst:
            used_cols.update(cols)
        if len(lst) > 1:
            sheet.warnings.append(f"연도 {y} 머리가 {len(lst)}번 나옵니다 — {y} 는 읽지 않습니다(원본 확인 필요)")
            continue
        cols = lst[0]
        if len(cols) == 1:
            periods.append((y, cols[0], []))
            continue
        finals = [c for c in cols if norm(status.get(c, "")).lower() in FINAL_TAGS]
        if len(finals) != 1:
            sheet.warnings.append(f"연도 {y} 아래 열이 {len(cols)}개인데 확정(Final) 열을 고를 수 없습니다 — {y} 는 읽지 않습니다")
            continue
        periods.append((y, finals[0], [c for c in cols if c != finals[0]]))
    periods.sort(key=lambda p: -p[0])
    sheet.periods = [y for y, _, _ in periods]

    for r in range(data_start, ws.max_row + 1):
        cell = ws.cell(r, label_col)
        v = cell.value
        if v is None or (isinstance(v, str) and not v.strip("  　\t")):
            _warn_orphan(sheet, ws, r, periods)
            continue
        if not isinstance(v, str) or norm(v).startswith("["):
            break
        rank, _ = split_prefix(v)
        row = ParsedRow(row_no=r, raw_label=v, label=display_label(v), rank=rank,
                        indent=int(cell.alignment.indent or 0) * 2 + leading_ws(v),
                        bold=bool(cell.font and cell.font.b), amounts={})
        for y, col, adj in periods:
            cv = ws.cell(r, col).value
            try:
                d, memo = parse_amount(cv)
            except AmountError as e:
                row.errors.append(f"{y}: {e}")
                d, memo = None, None
            row.amounts[y] = d
            row.raw_values[y] = raw_text(cv) if d is not None else None
            if memo == "float_noise":
                row.meta.setdefault("float_noise", []).append(y)
            extra = {status.get(c, get_column_letter(c)): ws.cell(r, c).value for c in adj
                     if ws.cell(r, c).value is not None and not isinstance(ws.cell(r, c).value, bool)}
            if extra:
                row.meta.setdefault("adjustments", {})[str(y)] = {k: str(x) for k, x in extra.items()}
        notes = {get_column_letter(c): ws.cell(r, c).value for c in range(1, ws.max_column + 1)
                 if c != label_col and c not in used_cols and isinstance(ws.cell(r, c).value, str)}
        if notes:
            row.meta["annotations"] = notes
        sheet.rows.append(row)
    return sheet


def _contiguous(status: dict[int, str], start: int, k: int) -> bool:
    return all(c in status for c in range(start + 1, k + 1))


def _warn_orphan(sheet: ParsedSheet, ws: Worksheet, r: int, periods) -> None:
    """라벨 없는 행에 선택 열 숫자가 있으면 경고 — 읽지는 않는다."""
    for y, col, _ in periods:
        v = ws.cell(r, col).value
        if isinstance(v, int | float) and not isinstance(v, bool):
            sheet.warnings.append(f"{r}행: 계정명 없이 {y} 금액이 있어 읽지 않았습니다 ({v})")
            return
