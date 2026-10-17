"""스코핑 엑셀 갱신 업로드 (2026-10-13, 13.9-118) — 회사가 정리한 스코핑 엑셀(원천 실무 양식)을 작성 중 스코핑에 반영한다.

원천 양식(`seeds/4__내부회계관리제도_Scoping.xlsx`, 템플릿의 출처)과 같은 열 배치를 읽는다 — 시트 4종
(2.1 BS · 2.2 PL · 2.3 주석 · 2.4 현금흐름), "계정과목" 머리글 행, 질적기준 10열, 질적 유의여부, 질적판단 근거,
유의성 결론. 열 위치는 머리글 문구로 찾는다(주석 시트는 한 칸 밀려 있다).

Regina 결정(STEP 0 Q1~Q6):
- Q1 맞추기: 표마다 ① 이름+구분명(번호·괄호 무시) ② 이름 ③ 템플릿 연결 이름 ④ 괄호·번호 뺀 이름 — 하나로 정해질 때만.
  남은 엑셀 행은 미리보기에서 사람이 고른다(`links` = {엑셀 키: 계정 id})
- Q2 반영: 질적 10요소 + 판단 근거. 결론은 **유의성 결론이 계산 결론(시스템 양적 + 엑셀 10요소)과 다를 때만** 수동
  판정(근거를 사유로) — 같으면 수동 판정을 비운다. 템플릿 시드와 같은 규칙
- Q3 금액은 반영하지 않는다(확정 재무제표가 원천) — 다르면 경고만
- Q4 바뀐 칸만 `edited`(화면 수정과 같은 `mark_edited`), 행위자는 업로드한 사람
- Q5 한쪽에만 있는 계정은 만들거나 지우지 않는다 — 목록으로 보인다. 해당 없음 계정은 건너뛴다
- Q6 내부회계 관리자, 작성 중일 때만(API 가 `_editable`)
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.scoping import (
    ORIGIN_TARGET_ACCOUNT,
    QUAL_FACTORS,
    RATING_ALIASES,
    STATEMENT_BS,
    STATEMENT_CF,
    STATEMENT_LABELS,
    STATEMENT_NOTE,
    STATEMENT_PL,
    STATEMENT_TYPES,
    Scoping,
    ScopingAccount,
    ScopingExcelAlias,
)
from app.services import scoping as svc
from app.services import scoping_calc as calc

# 시트 이름은 접두로 찾는다 — 원천은 끝에 공백이 붙은 이름도 있다("2.4 유의한 현금흐름 관련항목 ")
SHEET_PREFIX = {STATEMENT_BS: "2.1", STATEMENT_PL: "2.2", STATEMENT_NOTE: "2.3", STATEMENT_CF: "2.4"}
MANUAL_REASON_FALLBACK = "스코핑 엑셀 반영"
FIELD_LABELS = {"ratings": "질적 10요소", "qual_basis": "판단 근거", "manual": "수동 판정"}

_CELL_REF = re.compile(r"(?<![A-Za-z])\$?[A-Z]{1,3}\$?\d+")


class ExcelError(Exception):
    pass


@dataclass
class ExcelRow:
    statement_type: str
    row: int
    name: str
    group: str | None
    ratings: dict[str, str]
    qual_basis: str | None
    conclusion: str | None      # 유의성 결론(Y/N) — 없으면 None
    amount: int | None          # 당기 금액(경고용 — 반영하지 않는다)

    @property
    def key(self) -> str:
        return f"{self.statement_type}:{self.row}"


# ── 파싱 ────────────────────────────────────────────────

def _text(v) -> str | None:
    if v is None:
        return None
    s = " ".join(str(v).replace("\xa0", " ").split())
    return s or None


def _is_subtotal(formula) -> bool:
    """같은 시트 셀을 참조하는 수식이면 소계·합계(원천 시드와 같은 판정). 다른 시트 참조·숫자끼리 계산은 값이다."""
    return isinstance(formula, str) and formula.startswith("=") and "!" not in formula \
        and bool(_CELL_REF.search(formula[1:]))


def _header(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, 40):
        if any(_text(ws.cell(r, c).value) == "계정과목" for c in range(1, 10)):
            return r, {" ".join(str(ws.cell(r, c).value).split()): c
                       for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None}
    raise ExcelError(f"'{ws.title}' 시트에서 머리글(계정과목)을 찾지 못했습니다")


def _col(cols: dict[str, int], prefix: str) -> int | None:
    hits = [c for name, c in cols.items() if name.startswith(prefix)]
    return hits[0] if len(hits) == 1 else None


def _yn(v) -> str | None:
    s = (_text(v) or "").upper()
    return s if s in ("Y", "N") else None


def _amount(v) -> int | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, int | float | Decimal):
        return int(round(v))
    return None


def parse(content: bytes) -> tuple[dict[str, list[ExcelRow]], list[str]]:
    """엑셀 → 표별 계정 행. 오류(시트·머리글·평가값)는 ExcelError, 사소한 것은 경고 목록."""
    try:
        wbf = load_workbook(BytesIO(content))                  # 수식 — 소계 판정용
        wbv = load_workbook(BytesIO(content), data_only=True)  # 값
    except Exception as e:  # noqa: BLE001 — 열 수 없는 파일은 이유를 그대로 돌려준다
        raise ExcelError(f"엑셀 파일을 열 수 없습니다: {e}") from None
    out: dict[str, list[ExcelRow]] = {}
    warnings: list[str] = []
    for stype, prefix in SHEET_PREFIX.items():
        name = next((n for n in wbv.sheetnames if n.strip().startswith(prefix)), None)
        if name is None:
            raise ExcelError(f"'{prefix} …' 시트({STATEMENT_LABELS[stype]})가 없습니다")
        ws, wf = wbv[name], wbf[name]
        hdr, cols = _header(ws)
        q0, qual, basis, concl = (_col(cols, "질적기준"), _col(cols, "질적 유의"), _col(cols, "질적판단"),
                                  _col(cols, "유의성"))
        if None in (q0, basis, concl):
            raise ExcelError(f"'{name}' 시트의 열 위치(질적기준·질적판단 근거·유의성 결론)를 찾지 못했습니다")
        amt = next((c for n, c in cols.items() if re.match(r"FY\d{4}", n)), None)   # 주석 시트는 FY20XX — 금액 없음
        rows: list[ExcelRow] = []
        group = None
        for r in range(hdr + 1, ws.max_row + 1):
            label = _text(ws.cell(r, 2).value)
            if not label:
                continue
            c_val = _yn(ws.cell(r, concl).value)
            q_val = _yn(ws.cell(r, qual).value) if qual else None
            if (amt and _is_subtotal(wf.cell(r, amt).value)) or (c_val is None and q_val is None):
                group = label    # 소계·합계·구분 제목
                continue
            ratings: dict[str, str] = {}
            for i, f in enumerate(QUAL_FACTORS):
                v = _text(ws.cell(r, q0 + i).value)
                if v is None:
                    continue
                code = RATING_ALIASES.get(v.lower())
                if code is None:
                    raise ExcelError(f"{name}!{ws.cell(r, q0 + i).coordinate} 평가값 '{v}' 를 해석할 수 없습니다"
                                     " (High·Medium·Low)")
                ratings[f] = code
            if c_val is None:
                warnings.append(f"{STATEMENT_LABELS[stype]} {r}행 '{label}' 유의성 결론이 비어 있어 결론은 반영하지 않습니다")
            rows.append(ExcelRow(stype, r, label, group, ratings, _text(ws.cell(r, basis).value), c_val,
                                 _amount(ws.cell(r, amt).value) if amt else None))
        out[stype] = rows
    return out, warnings


# ── 맞추기 ──────────────────────────────────────────────

def _n(s: str | None) -> str:
    """비교 키 — 공백 제거, 원천의 회사 지칭('당사')을 템플릿 표기('회사')로."""
    return re.sub(r"\s+", "", (s or "").replace("\xa0", " ")).replace("당사", "회사")


def _core(s: str | None) -> str:
    """괄호 내용·앞 번호("6.", "(4)", "II.")를 뺀 키 — 대손충당금(매출채권) ↔ 대손충당금, (4) 이자의 수취 ↔ 이자의 수취."""
    k = _n(s)
    k = re.sub(r"^(\(\d+\)|\d+\.|[IVX]+\.)", "", k)
    return re.sub(r"\(.*?\)", "", k)


# 현금흐름 표기 변형(13.9-146) — 같은 항목을 해마다·작성자마다 다르게 쓴다. 비교용으로만 쓰고 저장하지 않는다
_CF_SYNONYMS = (("현금및현금성자산", "현금"), ("취득", "증가"), ("처분", "감소"))
_ALT = re.compile(r"^(.*?)([가-힣A-Za-z]+)\(([가-힣A-Za-z]+)\)(.*)$")


def _swap_alt(k: str) -> set[str]:
    """ "법인세의환급(납부)" → {법인세의환급, 법인세의납부}. 괄호 안 낱말은 **마지막 '의' 뒤**를 바꿔 끼운다."""
    m = _ALT.match(k)
    if not m:
        return {k}
    head, word, alt, tail = m.group(1), m.group(2), m.group(3), m.group(4)
    if "의" in word:
        cut = word.rindex("의") + 1
        head, word = head + word[:cut], word[cut:]
    return {head + word + tail, head + alt + tail}


def _cf_variants(s: str | None) -> set[str]:
    """현금흐름 항목 이름의 비교 키들 — 앞 번호를 떼고, "A(B)" 는 A·B 둘로 풀고, 같은 뜻 낱말을 하나로.

    "5. 법인세의 환급(납부)" → {법인세의환급, 법인세의납부}, "토지의 취득" → {토지의증가},
    "V. 기초의 현금" ↔ "기초의 현금및현금성자산" → {기초의현금}.
    """
    k = re.sub(r"^(\(\d+\)|\d+\.|[IVX]+\.)", "", _n(s))
    out = set()
    for v in _swap_alt(k):
        for a, b in _CF_SYNONYMS:
            v = v.replace(a, b)
        out.add(v)
    return out


def alias_key(statement_type: str, name: str | None, group: str | None) -> tuple[str, str, str]:
    """직접 연결 기억의 키 — 표 + 엑셀 이름 + 엑셀 구분(괄호·번호 제외)."""
    return statement_type, _n(name)[:200], _core(group)[:200]


@dataclass
class Pair:
    excel: ExcelRow
    account: ScopingAccount
    how: str


@dataclass
class Matching:
    pairs: list[Pair] = field(default_factory=list)
    excel_only: list[ExcelRow] = field(default_factory=list)
    account_only: list[ScopingAccount] = field(default_factory=list)


def match(excel: dict[str, list[ExcelRow]], accounts: list[ScopingAccount], template_names: dict[UUID, str],
          links: dict[str, UUID] | None = None,
          aliases: dict[tuple[str, str, str], tuple[str, str | None]] | None = None) -> Matching:
    """엑셀 행 ↔ 스코핑 계정. 규칙마다 **후보가 하나일 때만** 맞춘다(둘 이상이면 다음 규칙·사람에게).

    순서: 이번에 사람이 고른 연결 → **기억한 연결**(지난 업로드에서 고른 것, 이름 기준) → 이름 규칙들 →
    현금흐름 표기 변형. `aliases` = {alias_key: (계정 이름, 계정 구분)}.
    """
    links = links or {}
    aliases = aliases or {}
    m = Matching()
    used: set[UUID] = set()
    by_id = {a.id: a for a in accounts}
    for stype in STATEMENT_TYPES:
        pool = [a for a in accounts if a.statement_type == stype]
        rows = excel.get(stype, [])

        def take(row: ExcelRow, cands: list[ScopingAccount], how: str) -> bool:
            cands = [a for a in cands if a.id not in used]
            if len(cands) != 1:
                return False
            used.add(cands[0].id)
            m.pairs.append(Pair(row, cands[0], how))
            return True

        left: list[ExcelRow] = []
        # 사람이 고른 연결이 먼저 — 같은 표의 계정만
        for row in rows:
            aid = links.get(row.key)
            a = by_id.get(aid) if aid else None
            if a is not None and a.statement_type == stype and a.id not in used:
                used.add(a.id)
                m.pairs.append(Pair(row, a, "직접 연결"))
            else:
                left.append(row)
        # 기억한 연결 — 지난번에 사람이 고른 계정 이름(구분까지 기억했으면 구분도)으로 찾는다
        still = []
        for row in left:
            hit = aliases.get(alias_key(stype, row.name, row.group))
            cands = [a for a in pool if hit and _n(a.name) == _n(hit[0])
                     and (hit[1] is None or _core(a.group_label) == _core(hit[1]))] if hit else []
            if not take(row, cands, "기억한 연결"):
                still.append(row)
        left = still
        rules = (
            ("이름·구분", lambda r, a: _n(a.name) == _n(r.name) and _core(a.group_label) == _core(r.group)),
            ("이름", lambda r, a: _n(a.name) == _n(r.name)),
            ("템플릿 연결", lambda r, a: a.id in template_names and _n(template_names[a.id]) == _n(r.name)),
            ("괄호·번호 제외", lambda r, a: _core(a.name) == _core(r.name)),
        ) + ((("현금흐름 표기", lambda r, a: bool(_cf_variants(a.name) & _cf_variants(r.name))),)
             if stype == STATEMENT_CF else ())
        for how, rule in rules:
            still = []
            for row in left:
                if not take(row, [a for a in pool if rule(row, a)], how):
                    still.append(row)
            left = still
        m.excel_only.extend(left)
        m.account_only.extend(a for a in pool if a.id not in used)
    return m


# ── 차이 ────────────────────────────────────────────────

def _ratings_text(r: dict | None) -> str:
    r = r or {}
    return "".join(r.get(f) or "·" for f in QUAL_FACTORS)


def _template_names(db: Session, fs_account_ids: list[UUID]) -> dict[UUID, str]:
    """재무제표 계정 → 연결된 템플릿 계정 이름(8-C 확정 링크). 엑셀 계정명은 템플릿 이름을 따른다."""
    from app.models.financial_statement import FsTemplateLink
    from app.models.scoping import ScopingTemplateAccount
    if not fs_account_ids:
        return {}
    rows = db.query(FsTemplateLink.account_id, ScopingTemplateAccount.name).join(
        ScopingTemplateAccount, ScopingTemplateAccount.id == FsTemplateLink.template_account_id).filter(
        FsTemplateLink.account_id.in_(fs_account_ids), FsTemplateLink.is_deleted == False).all()  # noqa: E712
    return dict(rows)


def plan(db: Session, s: Scoping, excel: dict[str, list[ExcelRow]], links: dict[str, UUID] | None = None) -> dict:
    """미리보기 = 반영할 일. **저장하지 않는다.** 반영(`apply`)도 이 결과를 그대로 쓴다."""
    accounts = svc._active(db.query(ScopingAccount), ScopingAccount).filter(ScopingAccount.scoping_id == s.id).all()
    template_names = _template_names(db, [a.fs_account_id for a in accounts if a.fs_account_id])
    tnames = {a.id: template_names[a.fs_account_id] for a in accounts if a.fs_account_id in template_names}
    aliases = {(x.statement_type, x.excel_name_key, x.excel_group_key): (x.account_name, x.account_group)
               for x in svc._active(db.query(ScopingExcelAlias), ScopingExcelAlias).all()}
    m = match(excel, accounts, tnames, links, aliases)

    ev = {r["account"].id: r for r in svc.evaluate(db, s)["accounts"]}
    crit = svc.criteria(db, s)
    threshold, comparison = crit["qual_threshold"], crit["qual_comparison"]

    changes, skipped, warnings, no_quant = [], [], [], []
    for p in m.pairs:
        a, x = p.account, p.excel
        if a.not_applicable:
            skipped.append({"account_id": str(a.id), "account_name": a.name, "excel_name": x.name,
                            "statement_type": a.statement_type, "reason": "해당 없음 계정 — 반영하지 않습니다"})
            continue
        new_ratings = {f: x.ratings[f] for f in QUAL_FACTORS if f in x.ratings}
        new_basis = x.qual_basis
        manual, reason = a.manual_conclusion, a.manual_reason
        quant = ev[a.id]["quant"] if a.id in ev else None
        amount_differs = a.current_amount is not None and x.amount is not None and a.current_amount != x.amount
        if x.conclusion is not None and quant is None:
            # 금액이 없어 양적 판정을 못 하면 계산 결론도 없다 — 그때 결론을 수동 판정으로 넣으면 전부 수동이 된다
            no_quant.append(a.name)
        elif x.conclusion is not None and amount_differs:
            # 엑셀 결론은 엑셀 금액으로 낸 것 — 금액이 다르면(재분류 등) 그 결론으로 시스템 판정을 덮지 않는다(운영 실측:
            # 용역매출 엑셀 0 ↔ 재무제표 4.4억). 10요소·근거는 반영하고 결론은 사람이 본다 — 경고는 아래 금액 비교에서
            pass
        elif x.conclusion is not None:
            computed = calc.conclusion(quant, calc.qualitative(new_ratings, threshold, comparison))
            if x.conclusion != computed:
                manual, reason = x.conclusion, (new_basis or MANUAL_REASON_FALLBACK)
            else:
                manual, reason = None, None
        fields = []
        if svc.changed(a.ratings or {}, new_ratings):
            fields.append({"field": "ratings", "label": FIELD_LABELS["ratings"],
                           "before": _ratings_text(a.ratings), "after": _ratings_text(new_ratings)})
        if svc.changed(a.qual_basis, new_basis):
            fields.append({"field": "qual_basis", "label": FIELD_LABELS["qual_basis"],
                           "before": a.qual_basis, "after": new_basis})
        if svc.changed(a.manual_conclusion, manual) or (manual is not None and svc.changed(a.manual_reason, reason)):
            fields.append({"field": "manual", "label": FIELD_LABELS["manual"],
                           "before": a.manual_conclusion, "after": manual})
        if amount_differs:
            warnings.append(f"{STATEMENT_LABELS[a.statement_type]} '{a.name}' 금액이 엑셀({x.amount:,})과 "
                            f"스코핑({a.current_amount:,})이 다릅니다 — 금액은 확정 재무제표 값을 유지하고, "
                            "엑셀 결론은 반영하지 않습니다(화면에서 판정 확인)")
        if fields:
            changes.append({
                "key": x.key, "statement_type": a.statement_type, "account_id": str(a.id), "account_name": a.name,
                "excel_name": x.name, "excel_row": x.row, "how": p.how, "fields": fields,
                "_apply": {"ratings": new_ratings, "qual_basis": new_basis, "manual": manual, "reason": reason},
            })

    if no_quant:
        warnings.append(f"금액이 없어 양적 판정을 못 한 계정 {len(no_quant)}개는 결론(수동 판정)을 반영하지 않았습니다 "
                        f"— 예: {', '.join(no_quant[:5])}")
    n = len(changes)
    summary_text = f"계정 {n}건 변경" if n else "변경 없음"
    return {
        "matched": len(m.pairs), "changed_count": n, "summary_text": summary_text,
        # 반영 때 기억할 직접 연결(이번에 사람이 고른 것만) — 응답에서는 뺀다
        "_remember": [{"key": alias_key(p.excel.statement_type, p.excel.name, p.excel.group),
                       "account_name": p.account.name, "account_group": p.account.group_label}
                      for p in m.pairs if p.how == "직접 연결"],
        "remembered": sum(1 for p in m.pairs if p.how == "기억한 연결"),
        "changes": changes, "skipped": skipped, "warnings": warnings,
        "excel_only": [{"key": r.key, "statement_type": r.statement_type, "row": r.row, "name": r.name,
                        "group": r.group} for r in m.excel_only],
        "account_only": [{"id": str(a.id), "statement_type": a.statement_type, "name": a.name,
                          "group_label": a.group_label, "not_applicable": a.not_applicable}
                         for a in m.account_only],
        "by_statement": {st: {"excel": len(excel.get(st, [])),
                              "matched": sum(1 for p in m.pairs if p.account.statement_type == st),
                              "changed": sum(1 for c in changes if c["statement_type"] == st)}
                         for st in STATEMENT_TYPES},
    }


def remember(db: Session, items: list[dict]) -> int:
    """직접 연결을 기억한다(13.9-146) — 같은 키는 덮어쓴다(최근 선택이 이긴다). 커밋은 호출자."""
    for it in items:
        st, name_key, group_key = it["key"]
        row = db.query(ScopingExcelAlias).filter(ScopingExcelAlias.statement_type == st,
                                                 ScopingExcelAlias.excel_name_key == name_key,
                                                 ScopingExcelAlias.excel_group_key == group_key).first()
        if row is None:
            row = ScopingExcelAlias(statement_type=st, excel_name_key=name_key, excel_group_key=group_key,
                                    account_name=it["account_name"])
            db.add(row)
        row.is_deleted = False
        row.account_name, row.account_group = it["account_name"][:200], (it["account_group"] or None)
    return len(items)


def apply(db: Session, s: Scoping, result: dict) -> None:
    """미리보기의 할 일을 저장 — 화면 수정과 같은 규칙(바뀐 칸만 `edited`). 사람이 고른 연결은 기억한다. 커밋은 호출자."""
    remember(db, result.get("_remember", []))
    ids = [UUID(c["account_id"]) for c in result["changes"]]
    rows = {a.id: a for a in db.query(ScopingAccount).filter(ScopingAccount.id.in_(ids)).all()} if ids else {}
    for c in result["changes"]:
        a = rows[UUID(c["account_id"])]
        todo = c["_apply"]
        fields = {f["field"] for f in c["fields"]}
        if "ratings" in fields:
            old = a.ratings or {}
            for f in QUAL_FACTORS:
                if svc.changed(old.get(f), todo["ratings"].get(f)):
                    svc.mark_edited(db, s.id, ORIGIN_TARGET_ACCOUNT, a.id, f"ratings.{f}")
            a.ratings = dict(todo["ratings"])   # JSON 은 새 객체를 넣어야 변경이 감지된다
        if "qual_basis" in fields:
            a.qual_basis = todo["qual_basis"]
            svc.mark_edited(db, s.id, ORIGIN_TARGET_ACCOUNT, a.id, "qual_basis")
        if "manual" in fields:
            a.manual_conclusion, a.manual_reason = todo["manual"], todo["reason"]
            svc.mark_edited(db, s.id, ORIGIN_TARGET_ACCOUNT, a.id, "manual")


def public(result: dict) -> dict:
    """응답용 — 내부 할 일(`_apply`)을 뺀다."""
    out = {k: v for k, v in result.items() if k != "_remember"}
    out["changes"] = [{k: v for k, v in c.items() if k != "_apply"} for c in result["changes"]]
    return out


__all__ = ["ExcelError", "apply", "match", "parse", "plan", "public", "STATEMENT_BS", "STATEMENT_PL",
           "STATEMENT_CF", "STATEMENT_NOTE"]
