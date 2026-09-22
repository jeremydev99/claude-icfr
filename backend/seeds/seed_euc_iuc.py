"""원천 양식 2종 → EUC 파일·정보 항목 시드 (5-1, ADR-0033).

1차 출처는 레포의 원천 양식이다 — 파서가 원천 열을 빠뜨리는 일(평가주기 열, 2026-09-04)이
반복되지 않게 양식을 레포에 두고 여기서 직접 읽는다.

    backend/seeds/6__EUC_inventory.xlsx  — 시트 "EUC Inventory"  → euc_files
    backend/seeds/7__IUC_inventory.xlsx  — 시트 "IPE Template"   → information_items

실행 (컨테이너 내부, WORKDIR=/app):
    docker compose exec backend python -m seeds.seed_euc_iuc

**원천에 없는 값을 지어내지 않는다**(5-1 §2.3):
- **복잡도는 NULL(미평가).** 원천에는 `Macro 포함여부` 만 있다. "매크로·링크·모델" 단계는
  외부 링크·모델도 포함하므로 매크로가 없다는 것만으로 단계를 정할 수 없다.
  → 위험 등급도 산출되지 않아 "미평가" 로 보인다
- 원천의 `EUC Tool의 위험평가`(Low 등)는 `source_risk_rating` 에 **원천 참고값**으로 보존한다.
  산출 등급과 섞지 않는다 — 나중에 둘이 달라지면 그것이 검토 신호다
- 저장소 열의 "추후 사이냅소프트 작성" 은 값이 아니라 **빈칸 안내문**이다 → NULL

**읽는 방식**:
- `data_only=True` — EUC 시트 일부 셀이 **다른 워크북을 참조하는 수식**이다
  (`='[1]RCM 취합'!$Q$49`). 캐시된 값을 읽는다. 이 스크립트는 그 열(프로세스명·통제 상세)을
  쓰지 않지만, 쓰게 되면 원천 파일을 링크 없이 다시 저장했을 때 값이 비는지 먼저 확인할 것
- **데이터 행은 `No` 열이 숫자인 행만** — IUC 시트는 254행이지만 데이터는 8행이다(나머지는 빈 서식)
- 두 양식의 프로세스명 표기가 서로 다르다(`고정자산` vs `고정자산관리` 등). 프로세스는 통제에서
  파생되므로 **통제활동번호로만 잇는다**
- `식별된 EUC Control Test(예시)` 시트는 `#REF!` 가 섞인 예시 양식이라 시드 대상이 아니다

**연결** — IUC 의 `통제활동에 활용되는 정보` 와 EUC 의 `EUC 파일명` 이 같으면 같은 파일이다.
정보 Type 이 EUC 인 항목만 파일을 참조한다.

**재실행 안전** — 이미 있는 파일(이름)·항목(통제×이름)은 건너뛴다. 삭제 경로는 없다.
단일 트랜잭션이며 검증 불일치 시 전부 롤백한다.

⚠️ 정보 항목은 통제 id 를 FK 없이 가리킨다 — 이 시드 뒤에 `seed_baseline --reset` 을 돌리면
연결이 조용히 끊어진다(13.9-27).
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from app.core.database import SessionLocal
from app.core.tenant_context import DEFAULT_TENANT_CODE, reset_active_tenant, set_active_tenant
from app.models.euc import CHANGE_FREQUENCY_ALIASES, CHANGE_FREQUENCY_VALUES, RISK_GRADES, EucFile
from app.models.iuc import IMPORTANCE_VALUES, INFO_TYPE_ALIASES, INFO_TYPE_EUC, InformationItem
from app.models.tenant import Tenant
from app.services.control_resolver import resolve_controls

SEED_DIR = Path(__file__).resolve().parent
EUC_BOOK = SEED_DIR / "6__EUC_inventory.xlsx"
IUC_BOOK = SEED_DIR / "7__IUC_inventory.xlsx"
EUC_SHEET = "EUC Inventory"
IUC_SHEET = "IPE Template"

# 원천 빈칸 안내문 — 값으로 넣지 않는다
PLACEHOLDERS = {"추후 사이냅소프트 작성"}


def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return None if not s or s in PLACEHOLDERS or s.upper() == "N/A" else s


def _header_map(ws) -> tuple[int, dict[str, int]]:
    """`No` 로 시작하는 헤더 행을 찾아 {헤더: 열 번호}. 헤더 행 위치가 양식마다 다르다(3행·4행).

    헤더 안의 줄바꿈·연속 공백은 공백 하나로 접는다. **첫 줄만 키로 쓰지 않는다** — IUC 양식은
    `시스템/(줄바꿈)어플리케이션` 과 `시스템/(줄바꿈)어플리케이션 In-Scope…` 두 열의 첫 줄이
    같아 덮어써진다.
    """
    for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row and row[0] == "No":
            return idx, {" ".join(str(h).split()): i for i, h in enumerate(row) if h is not None}
    raise SystemExit(f"[중단] '{ws.title}' 에서 헤더 행(No)을 찾지 못했습니다")


def _data_rows(ws, header_idx: int):
    """`No` 열이 숫자인 행만 데이터다 — 빈 서식 행을 걸러낸다."""
    for row in ws.iter_rows(min_row=header_idx + 1, values_only=True):
        if row and isinstance(row[0], int | float):
            yield row


def _col(cols: dict[str, int], prefix: str) -> int:
    """헤더를 찾는다 — 정확히 같은 헤더가 있으면 그것, 없으면 **앞부분이 유일하게 맞는** 헤더.

    원천 헤더에 괄호 보조 문구가 붙어 있어 앞부분으로 찾지만, 두 열 이상이 걸리면 조용히 첫 열을
    고르지 않고 중단한다(`Report Logic` 과 `Report Logic 관련 control` 이 그런 쌍이다).
    """
    if prefix in cols:
        return cols[prefix]
    hits = [i for name, i in cols.items() if name.startswith(prefix)]
    if len(hits) == 1:
        return hits[0]
    raise SystemExit(f"[중단] 헤더 '{prefix}…' 가 {len(hits)}개 열에 걸립니다. 있는 헤더: {list(cols)}")


def _norm_frequency(v) -> str | None:
    s = _clean(v)
    if s is None:
        return None
    if s in CHANGE_FREQUENCY_VALUES:
        return s
    code = CHANGE_FREQUENCY_ALIASES.get(s.lower())
    if code is None:
        raise SystemExit(f"[중단] 파일변경주기 '{s}' 를 코드로 옮길 수 없습니다 — 상수에 별칭을 추가할 것")
    return code


def _norm_risk(v) -> str | None:
    """원천 위험평가 → 소문자 코드. 목록 밖 값은 그대로 보존한다(참고값이므로 버리지 않는다)."""
    s = _clean(v)
    if s is None:
        return None
    return s.lower() if s.lower() in RISK_GRADES else s


def _load() -> tuple[list[dict], list[dict]]:
    ws = load_workbook(EUC_BOOK, data_only=True)[EUC_SHEET]
    h, c = _header_map(ws)
    files = []
    for r in _data_rows(ws, h):
        macro = _clean(r[_col(c, "Macro")])
        files.append({
            "control_code": _clean(r[_col(c, "통제활동번호")]),
            "name": _clean(r[_col(c, "EUC 파일명")]),
            "description": _clean(r[_col(c, "EUC 파일설명")]),
            "has_macro": None if macro is None else macro.lower() in ("yes", "y"),
            "change_frequency": _norm_frequency(r[_col(c, "파일변경주기")]),
            "storage_path": _clean(r[_col(c, "EUC 파일 저장소")]),
            "managing_department": _clean(r[_col(c, "EUC 파일관리부서")]),
            "manager_name": _clean(r[_col(c, "EUC 파일관리자")]),
            "source_risk_rating": _norm_risk(r[_col(c, "EUC Tool의 위험평가")]),
            "source_risk_basis": _clean(r[_col(c, "위험평가근거")]),
        })

    ws = load_workbook(IUC_BOOK, data_only=True)[IUC_SHEET]
    h, c = _header_map(ws)
    items = []
    for r in _data_rows(ws, h):
        info_type_raw = _clean(r[_col(c, "정보 Type")])
        info_type = INFO_TYPE_ALIASES.get((info_type_raw or "").lower())
        if info_type is None:
            raise SystemExit(f"[중단] 정보 Type '{info_type_raw}' 를 코드로 옮길 수 없습니다")
        importance = _clean(r[_col(c, "Information 중요성")])
        if importance is not None and importance not in IMPORTANCE_VALUES:
            raise SystemExit(f"[중단] 중요성 '{importance}' 는 H/M/L 이 아닙니다")
        items.append({
            "control_code": _clean(r[_col(c, "통제활동번호")]),
            "name": _clean(r[_col(c, "통제활동에 활용되는 정보")]),
            "info_type": info_type,
            "importance": importance,
            "system_name": _clean(r[_col(c, "시스템/ 어플리케이션")]),
            "itgc_in_scope": _clean(r[_col(c, "시스템/ 어플리케이션 In-Scope")]),
            "source_data": _clean(r[_col(c, "기초정보 (Source Data)")]),
            "report_logic": _clean(r[_col(c, "Report Logic")]),
            "input_parameter": _clean(r[_col(c, "Input parameter (")]),
            "source_data_review": _clean(r[_col(c, "기초정보 신뢰성")]),
            "report_logic_control": _clean(r[_col(c, "Report Logic 관련")]),
            "input_parameter_review": _clean(r[_col(c, "Input parameter 정확성")]),
            "design_assessment_result": _clean(r[_col(c, "설계평가 결과")]),
        })
    return files, items


def seed(tenant_code: str = DEFAULT_TENANT_CODE) -> None:
    files, items = _load()
    print(f"  파싱: EUC 파일 {len(files)} / 정보 항목 {len(items)}")

    db = SessionLocal()
    tok = None
    try:
        tenant = db.query(Tenant).filter(Tenant.code == tenant_code).first()
        if tenant is None:
            raise SystemExit(f"[중단] 테넌트 코드 '{tenant_code}' 를 찾을 수 없습니다")
        tok = set_active_tenant(tenant.id)  # before_flush 가 tenant_id 를 stamp 한다(ADR-0025)

        controls = {c["code"]: c for c in resolve_controls(db)}
        missing = sorted({x["control_code"] for x in files + items} - set(controls))
        if missing:
            raise SystemExit(f"[중단] RCM 에 없는 통제활동번호: {missing} — baseline 시드를 먼저 확인할 것")

        existing_files = {
            f.name: f for f in db.query(EucFile).filter(EucFile.is_deleted == False).all()  # noqa: E712
        }
        created_files = 0
        for f in files:
            if f["name"] in existing_files:
                continue
            obj = EucFile(**{k: v for k, v in f.items() if k != "control_code"})
            db.add(obj)
            db.flush()
            existing_files[obj.name] = obj
            created_files += 1

        existing_items = {
            (i.control_id, i.name)
            for i in db.query(InformationItem).filter(InformationItem.is_deleted == False).all()  # noqa: E712
        }
        created_items = 0
        unlinked = []
        for it in items:
            control_id = controls[it["control_code"]]["id"]
            if (control_id, it["name"]) in existing_items:
                continue
            file = existing_files.get(it["name"]) if it["info_type"] == INFO_TYPE_EUC else None
            if it["info_type"] == INFO_TYPE_EUC and file is None:
                unlinked.append(it["name"])
            db.add(InformationItem(
                control_id=control_id, euc_file_id=file.id if file else None,
                **{k: v for k, v in it.items() if k != "control_code"},
            ))
            created_items += 1
        db.flush()

        if unlinked:
            raise SystemExit(f"[중단] EUC 정보인데 같은 이름의 EUC 파일이 없습니다: {unlinked}")

        db.commit()
        print(f"  생성: EUC 파일 {created_files} / 정보 항목 {created_items} "
              f"(이미 있어 건너뜀: 파일 {len(files) - created_files} / 항목 {len(items) - created_items})")
    except BaseException:
        db.rollback()
        raise
    finally:
        if tok is not None:
            reset_active_tenant(tok)
        db.close()


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="원천 양식 → EUC 파일·정보 항목 시드")
    ap.add_argument("--tenant", default=DEFAULT_TENANT_CODE, help="대상 테넌트 코드")
    seed(tenant_code=ap.parse_args().tenant)
