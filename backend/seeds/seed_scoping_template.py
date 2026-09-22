"""원천 실무 양식 → 전역 스코핑 템플릿 (6-1, ADR-0034 §2.7).

    backend/seeds/4__내부회계관리제도_Scoping.xlsx → scoping_templates(+texts, +accounts)

실행 (컨테이너 내부, WORKDIR=/app):
    docker compose exec backend python -m seeds.seed_scoping_template

**템플릿은 전역이다**(테넌트 인자가 없다). 제품 콘텐츠이며 스코핑 생성 시 테넌트로 복사된다.
이미 같은 `(code, version)` 이 있으면 아무것도 하지 않는다 — 개정은 version 을 올려 새로 넣는다.

**넣는 것**(ADR-0034 §2.7): 중요성 가이던스(적용기법 57·58·64~66), 질적 10요소 설명,
설정근거 문구, 계정 시트 4종의 판단 원칙과 각주, Notes 1~4, 벤치마크·설정율 기본값,
계정 191건(계정명·소속 구분·질적 10요소 평가값·판단 근거·수동 판정).

**넣지 않는 것**: 모든 금액(2022·2021), 표지(회사명), 주석 시트 FY2019 헤더,
업무프로세스·GITC 시트(6-2 범위 — 감사인 이메일이 여기 있다).

**문구 가공은 두 가지뿐이고, 전부 목록으로 출력한다**(무엇을 바꿨는지 보이게):
1. **회사 식별자 → "회사"** — 사이냅소프트·영문·법인격 표기·약칭, 감사인 법인명. `REPLACEMENTS`
2. **"초과" → "이상"** — 원천 설명문은 "2를 초과하면 중요", 수식은 `>=2` 로 서로 달랐다.
   실제 적용·검토된 것은 수식이다(ADR-0034 §2.4, 2026-09-22 결정)

적재 후 **금지 패턴 검사**를 한다 — 이메일(`@`)·회사명·감사인명이 하나라도 남으면 롤백하고 중단한다.

**계정과 소계는 수식 구조로 가른다**(계정명 문자열로 가르지 않는다):
- 금액 셀이 **같은 시트 셀을 참조하는 수식**(`=SUM(C21:C33)`, `=C20+C34`) → 소계
  (숫자끼리 계산 `=5005866000+1227635000` 이나 다른 시트 참조 `='2.1 …'!C34` 는 소계가 아니다)
- 결론 셀이 비었다 → 구분 제목(자산·부채 등). 금액 0 인 그룹 행(재고자산)도 여기 걸린다
- 나머지 → 계정. 소속 구분은 바로 앞의 소계·제목 이름

**수동 판정** — 결론 셀이 수식이 아니라 값으로 직접 입력됐고, **그 값이 계산 결론과 다른 행만**
수동 판정으로 넣는다(원천 6건). 같으면 수동 판정이 아니다(이익준비금).
"""
from __future__ import annotations

import argparse
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.models.scoping import (
    BENCHMARK_ADJUSTED_PBT,
    DEFAULT_QUAL_THRESHOLD,
    DEFAULT_TEMPLATE_CODE,
    QUAL_COMPARISON_GE,
    QUAL_FACTORS,
    RATING_ALIASES,
    STATEMENT_BS,
    STATEMENT_CF,
    STATEMENT_NOTE,
    STATEMENT_PL,
    ScopingTemplate,
    ScopingTemplateAccount,
    ScopingTemplateText,
)
from app.services import scoping_calc as calc
from app.services.scoping import default_criteria_payload

SOURCE = Path(__file__).resolve().parent / "4__내부회계관리제도_Scoping.xlsx"
TEMPLATE_VERSION = 1

SHEETS = {
    STATEMENT_BS: "2.1 유의한 계정과목(BS)",
    STATEMENT_PL: "2.2 유의한 계정과목(PL)",
    STATEMENT_NOTE: "2.3 유의한 주석",
    STATEMENT_CF: "2.4 유의한 현금흐름 관련항목 ",
}
MATERIALITY_SHEET = "1. 중요성 기준"
NOTES_SHEET = "Notes"

# 회사·감사인 식별자 → "회사". **긴 표기부터** 바꾼다(부분 치환으로 "주식회사 회사"가 되지 않게).
# 원천 적재 범위를 전부 훑은 결과(6-1 STEP 0): 회사명은 표지에만 있어 적재 범위에 없고,
# 감사인 법인명은 업무프로세스 시트 이메일에만 있어 역시 적재 범위 밖이다. 적재 범위에서
# 실제로 걸린 것은 주석 항목명의 "당사" 1건뿐이다. 목록은 다른 원천이 들어올 때를 위해 넓게 둔다.
REPLACEMENTS: list[tuple[str, str]] = [
    ("주식회사 사이냅소프트", "회사"),
    ("(주)사이냅소프트", "회사"),
    ("㈜사이냅소프트", "회사"),
    ("사이냅소프트", "회사"),
    ("Synapsoft", "회사"),
    ("SynapSoft", "회사"),
    ("SYNAPSOFT", "회사"),
    ("사이냅", "회사"),
    ("Synap", "회사"),
    ("당사", "회사"),
]
# 적재 후 하나라도 남으면 중단한다 — 이메일, 회사명, 감사인 법인명(대형 회계법인 영문·국문)
FORBIDDEN = re.compile(r"@|사이냅|synap|deloitte|딜로이트|안진|삼일|삼정|한영|pwc|kpmg|ernst|회계법인",
                       re.IGNORECASE)

# "초과" → "이상" (ADR-0034 §2.4). 원문 그대로 찾아 바꾼다 — 비슷한 문장이 있어도 건드리지 않는다
THRESHOLD_FIX = ("Rating결과가 2(Medium)을 초과할 경우", "Rating결과가 2(Medium) 이상일 경우")

_CELL_REF = re.compile(r"(?<![A-Za-z])\$?[A-Z]{1,3}\$?\d+")


class Loader:
    def __init__(self) -> None:
        self.changes: list[str] = []

    def clean(self, value, where: str) -> str | None:
        """문자열 정리 + 식별자 치환. 바꾼 것은 `changes` 에 원문→치환문으로 남긴다."""
        if value is None:
            return None
        s = str(value).strip()
        if not s:
            return None
        before = s
        if THRESHOLD_FIX[0] in s:
            s = s.replace(*THRESHOLD_FIX)
        for old, new in REPLACEMENTS:
            s = s.replace(old, new)
        if s != before:
            self.changes.append(f"{where}: {before!r} → {s!r}")
        return s


def _is_subtotal(amount) -> bool:
    """같은 시트 셀을 참조하는 수식이면 소계. 다른 시트 참조(`!`)나 숫자끼리 계산은 값이다."""
    return isinstance(amount, str) and amount.startswith("=") and "!" not in amount \
        and bool(_CELL_REF.search(amount[1:]))


def _header(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, 40):
        if any(ws.cell(r, c).value == "계정과목" for c in range(1, 10)):
            return r, {" ".join(str(ws.cell(r, c).value).split()): c
                       for c in range(1, ws.max_column + 1) if ws.cell(r, c).value}
    raise SystemExit(f"[중단] '{ws.title}' 에서 헤더(계정과목)를 찾지 못했습니다")


def _col(cols: dict[str, int], prefix: str) -> int | None:
    hits = [c for name, c in cols.items() if name.startswith(prefix)]
    return hits[0] if len(hits) == 1 else None


def _accounts(ld: Loader, wbf, wbv, stype: str) -> list[dict]:
    """계정 시트 → 계정 목록. 열 위치는 헤더 문구로 찾는다(주석 시트는 한 칸 밀려 있다)."""
    ws, wv = wbf[SHEETS[stype]], wbv[SHEETS[stype]]
    hdr, cols = _header(ws)
    q0 = _col(cols, "질적기준")
    concl = _col(cols, "유의성")
    basis = _col(cols, "질적판단")
    quant = next((c for n, c in cols.items() if n.startswith("양적") and "유의" in n), None)
    amount = next((c for c in range(1, ws.max_column + 1)
                   if re.match(r"FY\d{4}", str(wv.cell(hdr, c).value or ""))), None)
    if None in (q0, concl, basis, amount):
        raise SystemExit(f"[중단] '{ws.title}' 열 위치를 확정하지 못했습니다: {cols}")

    rows, group = [], None
    for r in range(hdr + 1, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if name is None or not str(name).strip():
            continue
        where = f"{ws.title}!B{r}"
        if _is_subtotal(ws.cell(r, amount).value) or ws.cell(r, concl).value is None:
            # 원천은 "유   동    자   산" 처럼 글자 사이를 띄워 썼다 — 표시용 공백이라 하나로 접는다
            group = " ".join((ld.clean(name, where) or "").split()) or None
            continue
        ratings = {}
        for i, f in enumerate(QUAL_FACTORS):
            v = ws.cell(r, q0 + i).value
            code = RATING_ALIASES.get(str(v).strip().lower()) if v is not None else None
            if v is not None and code is None:
                raise SystemExit(f"[중단] {ws.title}!{ws.cell(r, q0 + i).coordinate} 평가값 '{v}' 해석 불가")
            if code:
                ratings[f] = code

        # 수동 판정 — 결론이 값으로 입력됐고 계산 결론과 다른 행만
        concl_raw = ws.cell(r, concl).value
        manual = manual_reason = None
        if not (isinstance(concl_raw, str) and concl_raw.startswith("=")):
            quant_v = calc.NA if stype in (STATEMENT_NOTE, STATEMENT_CF) else wv.cell(r, quant).value
            qual_v = calc.qualitative(ratings, Decimal(DEFAULT_QUAL_THRESHOLD), QUAL_COMPARISON_GE)
            computed = calc.conclusion(quant_v, qual_v)
            if str(concl_raw).strip() != computed:
                manual = str(concl_raw).strip()
                manual_reason = ld.clean(ws.cell(r, basis).value, f"{ws.title}!{ws.cell(r, basis).coordinate}")
                if not manual_reason:
                    raise SystemExit(f"[중단] {where} 수동 판정인데 사유가 없습니다")

        rows.append({
            "statement_type": stype, "sort_order": len(rows) + 1, "group_label": group,
            "name": ld.clean(name, where), "ratings": ratings,
            "qual_basis": ld.clean(ws.cell(r, basis).value, f"{ws.title}!{ws.cell(r, basis).coordinate}"),
            "manual_conclusion": manual, "manual_reason": manual_reason,
        })
    return rows


def _texts(ld: Loader, wbv) -> list[dict]:
    """문구 → [{key, title, body}]. 키는 매뉴얼·화면이 참조하는 식별자라 바꾸지 않는다."""
    out: list[dict] = []
    m = wbv[MATERIALITY_SHEET]

    def block(ws, col: int, rows: range, where: str) -> str | None:
        lines = [ld.clean(ws.cell(r, col).value, f"{where}{r}") for r in rows]
        joined = "\n".join(x for x in lines if x)
        return joined or None

    out.append({"key": "materiality.guidance", "title": "중요성의 고려 (설계운영 적용기법 57·58·64~66)",
                "body": block(m, 4, range(6, 31), f"{MATERIALITY_SHEET}!D")})
    out.append({"key": "materiality.rationale", "title": "설정근거",
                "body": block(m, 4, range(66, 69), f"{MATERIALITY_SHEET}!D")})

    # 질적 10요소 — 제목 행(①~⑩)과 다음 제목 전까지의 설명
    circled = set("①②③④⑤⑥⑦⑧⑨⑩")
    title_rows = [r for r in range(72, m.max_row + 1)
                  if str(m.cell(r, 4).value or "").strip()[:1] in circled]
    if len(title_rows) != len(QUAL_FACTORS):
        raise SystemExit(f"[중단] 질적 요소 제목이 {len(title_rows)}개입니다(10개여야 함)")
    for i, (f, tr) in enumerate(zip(QUAL_FACTORS, title_rows)):
        end = title_rows[i + 1] if i + 1 < len(title_rows) else m.max_row + 1
        out.append({"key": f"qual.{f}", "title": ld.clean(m.cell(tr, 4).value, f"{MATERIALITY_SHEET}!D{tr}"),
                    "body": block(m, 4, range(tr + 1, end), f"{MATERIALITY_SHEET}!D")})

    for stype, sn in SHEETS.items():
        ws = wbv[sn]
        body = block(ws, 2, range(4, 11), f"{sn}!B")
        foot = next((ws.cell(r, 2).value for r in range(ws.max_row, 17, -1)
                     if str(ws.cell(r, 2).value or "").startswith("*")), None)
        if foot:
            body = f"{body}\n\n{ld.clean(foot, f'{sn}!각주')}"
        out.append({"key": f"statement.{stype}.principles", "title": "일반사항·질적중요성의 판단", "body": body})

    n = wbv[NOTES_SHEET]
    # "Note 1"~"Note 4" 만 — 시트 제목 "Notes"(B2)를 걸지 않게 숫자까지 본다
    starts = [r for r in range(1, n.max_row + 1)
              if re.match(r"Note \d", str(n.cell(r, 2).value or "").strip())]
    for i, r0 in enumerate(starts):
        end = starts[i + 1] if i + 1 < len(starts) else n.max_row + 1
        out.append({"key": f"notes.{i + 1}", "title": ld.clean(n.cell(r0, 3).value, f"Notes!C{r0}"),
                    "body": block(n, 3, range(r0 + 1, end), "Notes!C")})
    return out


def load_template(db: Session, source: Path = SOURCE, code: str = DEFAULT_TEMPLATE_CODE,
                  version: int = TEMPLATE_VERSION) -> tuple[ScopingTemplate, list[str], bool]:
    """템플릿을 적재한다. 반환: (템플릿, 치환 목록, 새로 만들었는가). 테스트도 이 함수를 쓴다."""
    existing = db.query(ScopingTemplate).filter(
        ScopingTemplate.code == code, ScopingTemplate.version == version,
        ScopingTemplate.is_deleted == False,  # noqa: E712
    ).first()
    if existing is not None:
        return existing, [], False

    wbf = load_workbook(source, data_only=False)
    wbv = load_workbook(source, data_only=True)
    ld = Loader()
    m = wbv[MATERIALITY_SHEET]
    rates = {BENCHMARK_ADJUSTED_PBT: str(m["F53"].value), "revenue": str(m["G53"].value)}

    tpl = ScopingTemplate(code=code, version=version, name="내부회계관리제도 스코핑 표준 템플릿",
                          source=source.name, default_benchmark=BENCHMARK_ADJUSTED_PBT,
                          default_rates=rates, default_smt_rate=str(m["G60"].value),
                          # 비율 가이드 범위·설정율 범위 기본값(원천 Note 1·2, 매출액은 비움 — 6-1b).
                          # 이미 적재된 템플릿은 마이그레이션 b1c2d3e4f5a6 이 같은 값으로 채운다
                          default_criteria=default_criteria_payload())
    db.add(tpl)
    db.flush()

    texts = _texts(ld, wbv)
    for i, t in enumerate(texts):
        if not t["body"]:
            raise SystemExit(f"[중단] 문구 '{t['key']}' 가 비었습니다")
        db.add(ScopingTemplateText(template_id=tpl.id, sort_order=i, **t))
    accounts = [a for st in SHEETS for a in _accounts(ld, wbf, wbv, st)]
    for a in accounts:
        db.add(ScopingTemplateAccount(template_id=tpl.id, **a))
    db.flush()

    # 금지 패턴 검사 — 적재한 모든 문자열
    leftovers = [
        v for v in [t["title"] for t in texts] + [t["body"] for t in texts]
        + [x for a in accounts for x in (a["name"], a["group_label"], a["qual_basis"], a["manual_reason"])]
        if v and FORBIDDEN.search(v)
    ]
    if leftovers:
        raise SystemExit(f"[중단] 식별자·이메일이 남았습니다: {leftovers[:5]}")
    return tpl, ld.changes, True


def main() -> None:
    db = SessionLocal()
    try:
        tpl, changes, created = load_template(db)
        if not created:
            print(f"  이미 있음: {tpl.code} v{tpl.version} — 아무것도 하지 않았습니다")
            return
        db.commit()
        n_acc = db.query(ScopingTemplateAccount).filter(ScopingTemplateAccount.template_id == tpl.id).count()
        n_txt = db.query(ScopingTemplateText).filter(ScopingTemplateText.template_id == tpl.id).count()
        n_man = db.query(ScopingTemplateAccount).filter(
            ScopingTemplateAccount.template_id == tpl.id,
            ScopingTemplateAccount.manual_conclusion.isnot(None)).count()
        print(f"  적재: {tpl.code} v{tpl.version} — 문구 {n_txt} / 계정 {n_acc} (수동 판정 {n_man})")
        print(f"  문구 가공 {len(changes)}건:")
        for c in changes:
            print(f"    {c}")
    except BaseException:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    argparse.ArgumentParser(description="원천 양식 → 전역 스코핑 템플릿").parse_args()
    main()
