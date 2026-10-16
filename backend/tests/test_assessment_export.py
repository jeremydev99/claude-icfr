"""운영평가 회차 엑셀 내보내기 (13.9-122) — 통제 1건 = 시트 1장, 증빙 이미지·PDF 페이지 붙이기, 내 담당만."""
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import minio_client
from app.api import evidence as evidence_api
from app.core.tenant_context import reset_active_tenant
from app.models.role_assignment import RoleAssignment
from tests.test_evidence_cycle import _control, _ctx, _cycle, _login, _user

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def storage(monkeypatch):
    store: dict[str, bytes] = {}

    class _Resp:
        def __init__(self, data: bytes):
            self._buf = io.BytesIO(data)

        def read(self):
            return self._buf.read()

        def stream(self, chunk_size: int = 8192):
            yield self._buf.read()

        def close(self):
            pass

        def release_conn(self):
            pass

    monkeypatch.setattr(evidence_api, "upload_object", lambda k, d, ct: store.__setitem__(k, d))
    monkeypatch.setattr(evidence_api, "get_object_stream", lambda k: _Resp(store[k]))
    monkeypatch.setattr(minio_client, "get_object_stream", lambda k: _Resp(store[k]))
    return store


def _png(w=400, h=200, color="red") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, format="PNG")
    return buf.getvalue()


def _pdf(pages: int) -> bytes:
    imgs = [Image.new("RGB", (300, 400), c) for c in (["white", "gray"] * pages)[:pages]]
    buf = io.BytesIO()
    imgs[0].save(buf, format="PDF", save_all=True, append_images=imgs[1:])
    return buf.getvalue()


def _media_count(content: bytes) -> int:
    return sum(1 for n in zipfile.ZipFile(io.BytesIO(content)).namelist() if n.startswith("xl/media/"))


def test_export_sheet_per_control_with_images_and_pdf_pages(client: TestClient, storage) -> None:
    db, tok = _ctx()
    try:
        owner = _user(db, "ex-owner@acme.example", "책임자X")
        other = _user(db, "ex-other@acme.example", "책임자Y")
        _user(db, "ex-mgr@acme.example", "관리자X", ("icfr_manager",))
        c1, c2 = _control(db, "X1"), _control(db, "X2")
        cyc = _cycle(db, "X1", c1)
        from app.models.assessment import CycleTarget
        db.add(CycleTarget(cycle_id=cyc, control_id=c2, control_code="EVX2-C"))
        db.add(RoleAssignment(scope="control", target_id=c1, role_name="control_owner", user_id=owner))
        db.add(RoleAssignment(scope="control", target_id=c2, role_name="control_owner", user_id=other))
        db.commit()
    finally:
        reset_active_tenant(tok)
        db.close()

    oh = _login(client, "ex-owner@acme.example", "pw123456")
    r = client.post("/api/evidence/files", headers=oh,
                    data={"cycle_id": str(cyc), "control_id": str(c1), "title": "급여대장 승인", "note": "전용남 CFO 승인"},
                    files={"file": ("캡처.png", _png(), "image/png")})
    assert r.status_code == 201, r.text
    assert r.json()["title"] == "급여대장 승인" and r.json()["note"] == "전용남 CFO 승인"
    assert client.post("/api/evidence/files", headers=oh, data={"cycle_id": str(cyc), "control_id": str(c1)},
                       files={"file": ("명세.pdf", _pdf(12), "application/pdf")}).status_code == 201

    # 회차 전체는 증빙 관리 역할만(13.9-137) — 통제책임자는 403, 관리자로 받는다
    assert client.get(f"/api/assessment/cycles/{cyc}/export", headers=oh).status_code == 403
    # 전체 — 통제 2건 = 시트 2장, 증빙 없는 통제도 시트가 있다
    res = client.get(f"/api/assessment/cycles/{cyc}/export", headers=_login(client, "ex-mgr@acme.example", "pw123456"))
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith(XLSX_MIME)
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["EVX1-C", "EVX2-C"]
    ws = wb["EVX1-C"]
    assert ws["A1"].value == "샘플수" and ws["B3"].value == "EVX1-P" and ws["F3"].value == "EVX1-R"
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert "급여대장 승인" in texts and "전용남 CFO 승인" in texts and "명세.pdf" in texts
    assert any("이후 2페이지는 원본 참조" in t for t in texts)          # PDF 12쪽 → 10쪽만 붙임
    assert _media_count(res.content) == 1 + 10                          # PNG 1 + PDF 페이지 10
    assert any("증빙 없음" in str(c.value) for row in wb["EVX2-C"].iter_rows() for c in row if c.value)

    # 내 담당만 — 통제책임자인 통제만
    mine = client.get(f"/api/assessment/cycles/{cyc}/export", headers=oh, params={"mine": "true"})
    assert mine.status_code == 200
    assert load_workbook(io.BytesIO(mine.content)).sheetnames == ["EVX1-C"]


def test_evidence_title_note_can_be_edited(client: TestClient, storage) -> None:
    db, tok = _ctx()
    try:
        owner = _user(db, "ex-owner2@acme.example", "책임자Z")
        c = _control(db, "X3")
        cyc = _cycle(db, "X3", c)
        db.add(RoleAssignment(scope="control", target_id=c, role_name="control_owner", user_id=owner))
        db.commit()
    finally:
        reset_active_tenant(tok)
        db.close()
    oh = _login(client, "ex-owner2@acme.example", "pw123456")
    fid = client.post("/api/evidence/files", headers=oh, data={"cycle_id": str(cyc), "control_id": str(c)},
                      files={"file": ("a.png", _png(), "image/png")}).json()["id"]
    r = client.patch(f"/api/evidence/files/{fid}", headers=oh, json={"title": "휴가신청", "note": "팀장 승인"})
    assert r.status_code == 200 and r.json()["title"] == "휴가신청"
    r = client.patch(f"/api/evidence/files/{fid}", headers=oh, json={"note": "  "})
    assert r.json()["note"] is None
