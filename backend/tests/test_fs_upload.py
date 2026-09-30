"""재무제표 엑셀 업로드 (8-B, ADR-0037 §3).

**실 샘플은 저장소에 넣지 않는다**(마스터 확정 Q1 — 내부 정산표·조정 내역 포함). 여기 워크북은
2025 감사 완료본에서 실측한 특징을 코드로 재현한 **합성 워크북**이다:
공시양식 = 기간당 2열·NBSP 띄어쓰기 합계·들여쓰기·PL 사다리·CF 차감 항목·검산 열·주당이익,
가로 연도형 = 연도 머리 + Final/조정 열·매핑 열·굵게 계층·원본 소계 수식 누락·연도 중복·단위 없음.
실 샘플 회귀는 `test_fs_upload_real_sample.py`(로컬 전용, `FS_SAMPLE_FILE` 있을 때만).

업로드는 HTTP 로 부른다(13.9-35 교훈). 계정 마스터가 테넌트 단위라 **테스트마다 새 테넌트**를 쓴다.
"""
import json
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import func, select

from app.core.security import hash_password
from app.core.tenant_context import reset_active_tenant, set_active_tenant
from app.models.financial_statement import FsAccount, FsAmount, FsStatement
from app.models.tenant import Tenant, UserTenantAccess
from app.models.user import User
from app.models.user_mgmt import UserRole
from app.services import financial_statement as svc_fs
from app.services import fs_upload
from app.services.fs_upload import structure
from app.services.fs_upload.cells import AmountError, display_label, parse_amount, split_prefix
from tests.conftest import TestingSessionLocal

PW = "pw123456"
D = Decimal
NB = " "
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── 테넌트·사용자 ─────────────────────────────────────────────

def _tenant(client: TestClient, roles: tuple[str, ...] = ("icfr_manager",)) -> tuple[dict, object]:
    """새 테넌트 + 그 테넌트 전용 사용자. (헤더, tenant_id)."""
    tid = uuid4()
    email = f"fsu-{tid.hex[:8]}@acme.example"
    db = TestingSessionLocal()
    try:
        db.add(Tenant(id=tid, name=f"업로드{tid.hex[:4]}", code=f"FSU-{tid.hex[:6]}", is_active=True))
        u = User(email=email, hashed_password=hash_password(PW), display_name="fsu", role="user", is_active=True)
        db.add(u)
        db.commit()
        db.add(UserTenantAccess(user_id=u.id, tenant_id=tid, role="user"))
        db.commit()
        tok = set_active_tenant(tid)
        try:
            for r in roles:
                db.add(UserRole(user_id=u.id, role_name=r))
            db.commit()
        finally:
            reset_active_tenant(tok)
    finally:
        db.close()
    r = client.post("/api/auth/login", data={"username": email, "password": PW})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"], "X-Tenant-Id": str(tid)}, tid


def _count(tid, model) -> int:
    db = TestingSessionLocal()
    tok = set_active_tenant(tid)
    try:
        return db.scalar(select(func.count()).select_from(model).where(model.is_deleted == False))  # noqa: E712
    finally:
        reset_active_tenant(tok)
        db.close()


def _post(client, headers, wb: Workbook, name="fs.xlsx", **form):
    buf = BytesIO()
    wb.save(buf)
    data = {k: (json.dumps(v) if isinstance(v, dict) else str(v)) for k, v in form.items() if v is not None}
    return client.post("/api/fs/upload", headers=headers, data=data,
                       files={"file": (name, buf.getvalue(), XLSX)})


def _body(r) -> dict:
    j = r.json()
    return j["detail"] if "detail" in j and isinstance(j["detail"], dict) else j


def _rows(body: dict) -> dict[str, dict]:
    return {r["label"]: r for r in body["rows"]}


# ── 합성 워크북: 공시양식형 ───────────────────────────────────

def _disclosure(ws, title: str, rows: list[tuple], *, unit="(단위 : 원)", cur_col=3,
                terms=((26, 2025), (25, 2024)), check_col=True) -> None:
    """rows: (라벨, 들여쓰기, 당기, 전기, 바깥열?). 기간당 2열(안쪽·바깥), 오른쪽 검산 열."""
    ws["A1"] = title
    ws["A1"].font = Font(bold=True)
    for i, (term, year) in enumerate(terms):
        ws.cell(2 + i, 1, f"제{term}기 {year}년 12월 31일 현재")
    ws["A5"] = "주식회사 합성"
    ws.cell(5, cur_col + 3, unit)
    ws.cell(6, 1, "과목")
    ws.cell(6, 2, "주석") if cur_col > 2 else None
    for i, (term, _y) in enumerate(terms):
        ws.cell(6, cur_col + 2 * i, f"제 {term} 기말")
        ws.merge_cells(start_row=6, start_column=cur_col + 2 * i, end_row=6, end_column=cur_col + 2 * i + 1)
    for n, (label, indent, cur, prior, outer) in enumerate(rows, start=7):
        c = ws.cell(n, 1, label)
        if indent:
            c.alignment = Alignment(indent=indent)
        for i, v in enumerate((cur, prior)):
            if v is not None:
                ws.cell(n, cur_col + 2 * i + (1 if outer else 0), v)
        if check_col and (cur is not None):
            ws.cell(n, cur_col + 5, True)          # 검산 열 — 읽으면 안 된다


BS_ROWS = [
    ("자산", 0, None, None, False),
    ("I. 유동자산", 0, 300, 250, True),
    ("현금및현금성자산", 1, 100, 100, False),
    ("매출채권", 1, 200, 150, False),
    ("II. 비유동자산", 0, 500, 450, True),
    ("유형자산", 1, 500, 450, False),
    (f"자 {NB} 산 {NB} 총 {NB} 계", 0, 800, 700, True),
    ("부채", 0, None, None, False),
    ("I. 유동부채", 0, 120, 100, True),
    ("매입채무", 1, 120, 100, False),
    (f"부 {NB} 채 {NB} 총 {NB} 계", 0, 120, 100, True),     # 자식이 하나인 합계
    ("자본", 0, None, None, False),
    ("I. 자본금", 0, 500, 500, False),
    ("II. 이익잉여금", 0, 180, 100, False),
    (f"자 {NB} 본 {NB} 총 {NB} 계", 0, 680, 600, True),
    ("부 채 및 자 본 총 계", 0, 800, 700, True),
]

PL_ROWS = [
    ("I. 영업수익", 0, 1000, 900, True),
    ("II. 영업비용", 0, 600, 500, True),
    ("III. 영업이익", 0, 400, 400, True),
    ("IV. 금융손익", 0, 30, -10, True),
    ("1. 금융수익", 1, 50, 20, False),
    ("2. 금융비용", 1, 20, 30, False),
    ("V. 법인세차감전순이익", 0, 430, 390, True),
    ("VI. 법인세비용", 0, 30, 20, True),
    ("VII. 당기순이익", 0, 400, 370, True),
    ("VIII. 주당손익", 0, None, None, False),
    (" 기본주당순이익", 0, 1181, 1000, True),
]

CF_ROWS = [
    ("I. 영업활동으로 인한 현금흐름", 0, 150, 100, True),
    ("1. 영업활동에서 창출된 현금", 0, 150, 100, False),     # 소계가 안쪽 열에 있다(실측 CF)
    (f"(1){NB}당기순이익", 0, 100, 80, False),
    ("(2) 비용가산 :", 0, 70, 40, False),
    ("감가상각비", 1, 70, 40, False),
    ("(3) 수익차감 :", 0, 20, 20, False),                  # 양수 표시 후 차감
    ("이자수익", 1, 20, 20, False),
    ("II. 투자활동으로 인한 현금흐름", 0, -50, -30, True),
    ("1. 투자활동으로 인한 현금유입액", 0, 10, 0, False),
    ("단기금융상품의 감소", 1, 10, 0, False),
    ("2. 투자활동으로 인한 현금유출액", 0, 60, 30, False),
    ("비품의 취득", 1, 60, 30, False),
    ("III. 현금및현금성자산의 증가", 0, 100, 70, True),
    ("IV. 기초의 현금및현금성자산", 0, 200, 130, True),
    ("V. 기말의 현금및현금성자산", 0, 300, 200, True),
]


def _wb(*sheets: tuple[str, str, list[tuple]], **kw) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for name, title, rows in sheets:
        _disclosure(wb.create_sheet(name), title, rows, **kw)
    return wb


def bs_wb(**kw) -> Workbook:
    return _wb(("BS공시", "재 무 상 태 표", BS_ROWS), **kw)


# ── 합성 워크북: 가로 연도형 ──────────────────────────────────

def horizontal_bs(*, stale_parent: bool = True, unit: str | None = None, orphan: bool = False,
                  cash_2025=120) -> Workbook:
    """1행 연도(2023·2024·2025), 2025 는 조정전/IFRS조정/감사수정/Final 4열 + 매핑 열."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BS정산표"
    ws.append(["과목", 2023, 2024, 2025])
    ws.append([None, "Final", "Final", "조정전", "IFRS조정", "감사수정", "Final"])
    if unit:
        ws["J1"] = unit
    rows = [  # (라벨, 굵게, 2023, 2024, 2025 Final, 조정전/IFRS조정, 매핑)
        ("자산", True, None, None, None, None, None),
        ("Ⅰ.유동자산", True, 145, 174, 183, None, None),
        ("(1)당좌자산", True, 145, 174, 183, None, None),
        ("현금", False, 100, 110, cash_2025, (125, -5), "현금및현금성자산"),
        ("매출채권", False, 50, 70, 70, None, "매출채권및기타채권"),
        ("대손충당금", False, -5, -6, -7, None, "매출채권및기타채권"),
        ("Ⅱ.비유동자산", True, 300, 300, 300, None, None),
        ("토지", False, 300, 300, 300, None, "유형자산"),
        ("자산총계", True, 445, 474, 483, None, None),
        ("부채", True, None, None, None, None, None),
        ("Ⅰ.유동부채", True, 45, 74, 83, None, None),
        ("미지급금", False, 45, 64, 83, None, "매입채무및기타채무"),
        ("유동성장기부채", False, 0, 10, 0, None, "단기차입금"),
        ("유동성전환사채", False, 0, 0, 0, None, "단기차입금"),
        ("상환할증금", False, 0, 0, 0, None, "단기차입금"),     # 2025 만 보면 0 = 0 + 0
        ("부채총계", True, 45, 74, 83, None, None),
        ("자본", True, None, None, None, None, None),
    ]
    if stale_parent:   # 원본 수식 누락 — 자본조정을 빼고 더한 "지배주주지분"(실측 BS정산표)
        rows.append(("지배주주의소유주에게귀속되는지분", True, 500, 500, 500, None, None))
    rows += [
        ("Ⅰ.자본금", True, 500, 500, 500, None, None),
        ("자본금", False, 500, 500, 500, None, "I. 자본금"),
        ("Ⅱ.자본조정", True, -100, -100, -100, None, None),
        ("자기주식", False, -100, -100, -100, None, "II. 기타불입자본"),
        ("자본총계", True, 400, 400, 400, None, None),
        ("부채및자본총계", True, 445, 474, 483, None, None),
    ]
    for label, bold, a, b, fin, adj, mapto in rows:
        ws.append([label, a, b, adj[0] if adj else fin, adj[1] if adj else None, None, fin, None, mapto])
        ws.cell(ws.max_row, 1).font = Font(bold=bold)
    if orphan:
        ws.append([None, None, None, None, None, None, 7])       # 라벨 없이 Final 금액
    ws.append([None, True, True, True, True, True, True])          # 검산 행
    ws.append([1, 2, 3, 4, 5, 6, 7])                               # 행번호 행 — 여기서 멈춘다
    ws.append(["[영업외손익분류]"])
    ws.append(["금융이익", 999, 999, 999, None, None, 999])
    return wb


def horizontal_pl_dup_year() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "PL정산표"
    ws.append(["과목", 2022, 2022, 2023, 2024])       # 원본 결함 — 2022 두 번
    ws.append([None, "Final", "Final", "Final", "Final"])
    for label, bold, *v in [
        ("I.영업수익", True, 100, 110, 120, 130),
        ("매출", False, 100, 110, 120, 130),
        ("II.영업비용", True, 60, 60, 70, 80),
        ("급여", False, 40, 40, 50, 60),
        ("감가상각비", False, 20, 20, 20, 20),
        ("III.영업이익", True, 40, 50, 50, 50),
    ]:
        ws.append([label, *v])
        ws.cell(ws.max_row, 1).font = Font(bold=bold)
    return wb


# ── 1. 셀 해석 ────────────────────────────────────────────────

def test_parse_amount_rules() -> None:
    assert parse_amount("(1,234)") == (D(-1234), "text")
    assert parse_amount("△500") == (D(-500), "text")
    assert parse_amount("-") == (D(0), "dash")
    assert parse_amount(None) == (None, None) and parse_amount("  ") == (None, None)
    assert parse_amount(True) == (None, None)                          # 검산 열
    assert parse_amount(120.0000001) == (D(120), "float_noise")
    for bad in (12.5, "1,234.5", "abc"):
        with pytest.raises(AmountError):
            parse_amount(bad)


def test_labels_and_prefixes() -> None:
    assert split_prefix("III. 영업이익") == (1, "영업이익")
    assert split_prefix("Ⅱ.비유동자산") == (1, "비유동자산")
    assert split_prefix("v. 기타영업외손익")[0] == 1
    assert split_prefix("1. 금융수익") == (2, "금융수익")
    assert split_prefix("(3) 수익차감 :") == (3, "수익차감 :")
    assert split_prefix("가. 매출") == (4, "매출")
    assert split_prefix("현금") == (0, "현금")
    assert display_label(f"자 {NB} 산 {NB} 총 {NB} 계") == "자산총계"
    assert display_label("  기타포괄손익 평가손익") == "기타포괄손익 평가손익"


# ── 2. 구조 추론 (DB 없음) ────────────────────────────────────

def test_disclosure_bs_tree_sections_and_both_periods() -> None:
    s = fs_upload.parse_sheet(bs_wb()["BS공시"])
    assert (s.kind, s.statement_type, s.unit, s.periods) == ("disclosure_form", "BS", 1, [2025, 2024])
    rows = {r.label: r for r in s.rows}
    assert rows["자산"].kind == "section_header" and not rows["자산"].is_account
    assert rows["자산총계"].kind == "footer"
    assert [c.label for c in rows["자산총계"].children] == ["유동자산", "비유동자산"]
    assert rows["유동자산"].kind == "header" and rows["현금및현금성자산"].parent is rows["유동자산"]
    assert [c.label for c in rows["부채총계"].children] == ["유동부채"]         # 자식 하나인 합계
    assert [c.label for c in rows["부채및자본총계"].children] == ["부채총계", "자본총계"]
    assert rows["부채및자본총계"].section == "liability_equity"
    assert rows["매입채무"].section == "liability" and rows["이익잉여금"].section == "equity"
    assert rows["현금및현금성자산"].meta["col"] == {"2025": "inner", "2024": "inner"}
    assert structure.subtotal_diffs(s) == []


def test_disclosure_pl_ladder_signs_sections_and_eps_excluded() -> None:
    s = fs_upload.parse_sheet(_wb(("PL공시", "포 괄 손 익 계 산 서", PL_ROWS))["PL공시"])
    rows = {r.label: r for r in s.rows}
    op = rows["영업이익"]
    assert op.kind == "footer" and [(c.label, c.sign) for c in op.children] == [("영업수익", 1), ("영업비용", -1)]
    assert [(c.label, c.sign) for c in rows["금융손익"].children] == [("금융수익", 1), ("금융비용", -1)]
    assert [c.label for c in rows["법인세차감전순이익"].children] == ["영업이익", "금융손익"]
    assert [(c.label, c.sign) for c in rows["당기순이익"].children] == [("법인세차감전순이익", 1), ("법인세비용", -1)]
    assert (rows["영업수익"].section, rows["영업비용"].section, op.section) == ("revenue", "expense", "profit")
    assert rows["금융비용"].section == "expense" and rows["금융손익"].section == "profit"
    assert rows["주당손익"].excluded and rows["기본주당순이익"].excluded
    assert "excluded_per_share" in rows["주당손익"].flags
    assert structure.subtotal_diffs(s) == []


def test_disclosure_cf_deductions_and_inner_column_subtotal() -> None:
    s = fs_upload.parse_sheet(_wb(("CF공시", "현 금 흐 름 표", CF_ROWS), cur_col=2)["CF공시"])
    rows = {r.label: r for r in s.rows}
    gen = rows["영업활동에서 창출된 현금"]
    assert gen.kind == "header" and gen.meta["col"]["2025"] == "inner"
    assert [(c.label, c.sign) for c in gen.children] == [("당기순이익", 1), ("비용가산 :", 1), ("수익차감 :", -1)]
    assert [c.sign for c in rows["투자활동으로 인한 현금흐름"].children] == [1, -1]
    assert [c.label for c in rows["현금및현금성자산의 증가"].children] == \
        ["영업활동으로 인한 현금흐름", "투자활동으로 인한 현금흐름"]
    assert [c.label for c in rows["기말의 현금및현금성자산"].children] == \
        ["현금및현금성자산의 증가", "기초의 현금및현금성자산"]
    assert rows["이자수익"].section == "operating" and rows["비품의 취득"].section == "investing"
    assert rows["기말의 현금및현금성자산"].section == "cf_other"
    assert structure.subtotal_diffs(s) == []


def test_horizontal_final_column_adjustments_annotations_and_stop() -> None:
    s = fs_upload.parse_sheet(horizontal_bs(stale_parent=False)["BS정산표"])
    assert (s.kind, s.statement_type, s.unit, s.periods) == ("horizontal_years", "BS", None, [2025, 2024, 2023])
    rows = {r.label: r for r in s.rows}
    cash = rows["현금"]
    assert cash.amounts == {2025: D(120), 2024: D(110), 2023: D(100)}           # Final 열
    assert cash.meta["adjustments"]["2025"] == {"조정전": "125", "IFRS조정": "-5"}
    assert cash.meta["annotations"] == {"I": "현금및현금성자산"}
    assert "금융이익" not in rows                                                # 행번호 행에서 멈춤
    assert rows["당좌자산"].parent is rows["유동자산"] and cash.parent is rows["당좌자산"]
    # 2025 만 보면 0 = 0 + 0 이지만 2024 가 맞지 않으므로 합계로 보지 않는다
    assert rows["상환할증금"].kind == "leaf" and rows["상환할증금"].parent is rows["유동부채"]
    assert structure.subtotal_diffs(s) == []


def test_horizontal_stale_source_subtotal_is_reported() -> None:
    s = fs_upload.parse_sheet(horizontal_bs(stale_parent=True)["BS정산표"])
    rows = {r.label: r for r in s.rows}
    assert "sign_unresolved" in rows["지배주주의소유주에게귀속되는지분"].flags
    assert "footer_mismatch" in rows["자본총계"].flags
    diffs = {(d["label"], d["fiscal_year"]) for d in structure.subtotal_diffs(s)}
    assert ("지배주주의소유주에게귀속되는지분", 2025) in diffs and ("자본총계", 2023) in diffs


def test_horizontal_duplicate_year_is_skipped_with_warning() -> None:
    s = fs_upload.parse_sheet(horizontal_pl_dup_year()["PL정산표"])
    assert s.periods == [2024, 2023]
    assert any("2022" in w for w in s.warnings)
    rows = {r.label: r for r in s.rows}
    assert [(c.label, c.sign) for c in rows["영업이익"].children] == [("영업수익", 1), ("영업비용", -1)]


def test_solve_signs_needs_all_periods_and_evidence() -> None:
    assert structure.solve_signs([D(0), D(10)], [[D(0), D(10)], [D(0), D(0)]]) == ([1, 1], False)
    assert structure.solve_signs([D(0), D(0)], [[D(0), D(10)], [D(0), D(0)]]) is None
    assert not structure.has_evidence([D(0)], [[D(0)], [D(0)]])


def test_unsupported_and_not_statement_sheets() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "자 본 변 동 표"
    ws["A6"] = "과목"
    ws["B6"] = "자본금"
    s = fs_upload.parse_sheet(ws)
    assert s.statement_type == "SCE" and "자본변동표" in s.errors[0]
    other = Workbook().active
    other.append(["거래처", "금액"])
    with pytest.raises(ValueError):
        fs_upload.parse_sheet(other)


# ── 3. 업로드 API ─────────────────────────────────────────────

def test_preview_does_not_write_and_matches_commit(client: TestClient) -> None:
    h, tid = _tenant(client)
    pre = _post(client, h, bs_wb(), mode="preview")
    assert pre.status_code == 200, pre.text
    b = pre.json()
    assert b["can_commit"] and not b["committed"] and b["master_empty"]
    assert (b["sheet"], b["unit"], b["fiscal_years"]) == ("BS공시", 1, [2025])
    assert [s["ok"] for s in b["statements"]] == [True] and b["statements"][0]["finalize_candidate"]
    assert _count(tid, FsAccount) == 0 and _count(tid, FsStatement) == 0       # 저장 안 함

    com = _post(client, h, bs_wb(), mode="commit")
    assert com.status_code == 200, com.text
    c = com.json()
    st = c["statements"][0]
    assert c["committed"] and st["status"] == "final" and st["finalized"]
    assert st["checks_count"] == b["statements"][0]["checks_count"]
    assert _count(tid, FsAccount) == 13 and _count(tid, FsAmount) == 13        # 섹션 머리 3행 제외
    # 확정 후 조회 API 의 검증도 같다
    v = client.get(f"/api/fs/statements/{st['statement_id']}/validation", headers=h).json()
    assert v["ok"] and len(v["checks"]) == st["checks_count"]


def test_commit_raw_columns_preserved(client: TestClient) -> None:
    h, tid = _tenant(client)
    c = _post(client, h, horizontal_bs(stale_parent=False), mode="commit", unit=1).json()
    sid = next(s["statement_id"] for s in c["statements"] if s["fiscal_year"] == 2025)
    d = client.get(f"/api/fs/statements/{sid}", headers=h).json()
    assert (d["source_kind"], d["source_filename"], d["source_sheet"]) == ("horizontal_years", "fs.xlsx", "BS정산표")

    def find(nodes, name):
        for n in nodes:
            if n["name"] == name:
                return n
            got = find(n["children"], name)
            if got:
                return got
    cash = find(d["tree"], "현금")
    assert D(cash["amount"]) == D(120) and cash["raw_label"] == "현금" and cash["raw_row_no"] == 6
    assert cash["raw_meta"]["adjustments"] == {"조정전": "125", "IFRS조정": "-5"}
    assert cash["raw_meta"]["annotations"] == {"I": "현금및현금성자산"}
    assert D(find(d["tree"], "대손충당금")["amount"]) == D(-7)


def test_horizontal_latest_year_final_only_when_valid(client: TestClient) -> None:
    h, _ = _tenant(client)
    ok = _post(client, h, horizontal_bs(stale_parent=False), mode="commit", unit=1).json()
    assert {s["fiscal_year"]: s["status"] for s in ok["statements"]} == {2025: "final", 2024: "draft", 2023: "draft"}

    h2, _ = _tenant(client)
    bad = _post(client, h2, horizontal_bs(stale_parent=True), mode="commit", unit=1)
    assert bad.status_code == 200, bad.text                          # 원본 오류가 있어도 draft 저장(Q8)
    sts = bad.json()["statements"]
    assert {s["fiscal_year"]: s["status"] for s in sts} == {2025: "draft", 2024: "draft", 2023: "draft"}
    latest = next(s for s in sts if s["fiscal_year"] == 2025)
    assert not latest["ok"] and latest["finalize_candidate"] and not latest["finalized"]
    # 원본 소계 불일치는 임시계정(원본 차이)으로 받고, 미해결 임시계정이 확정을 막는다(마스터 지시 2026-09-30)
    assert {e["rule"] for e in latest["errors"]} == {"suspense_unresolved"}
    assert {(x["parent_name"], D(x["amount"])) for x in latest["suspense"]} == \
        {("지배주주의소유주에게귀속되는지분", D(100)), ("자본총계", D(-100))}

    h4, _ = _tenant(client)
    off = _post(client, h4, horizontal_bs(stale_parent=True), mode="commit", unit=1, suspense="false").json()
    latest = next(s for s in off["statements"] if s["fiscal_year"] == 2025)
    assert {e["rule"] for e in latest["errors"]} == {"subtotal"} and latest["suspense"] == []
    assert all(e["raw_row_no"] for e in latest["errors"])

    h3, _ = _tenant(client)
    off = _post(client, h3, horizontal_bs(stale_parent=False), mode="commit", unit=1, finalize="false").json()
    assert {s["status"] for s in off["statements"]} == {"draft"}


def test_unit_required_and_must_match(client: TestClient) -> None:
    h, tid = _tenant(client)
    pre = _post(client, h, horizontal_bs(), mode="preview").json()
    assert not pre["can_commit"] and any("단위" in e for e in pre["errors"])
    assert _post(client, h, horizontal_bs(), mode="commit").status_code == 422
    assert _post(client, h, horizontal_bs(unit="(단위: 백만원)"), mode="preview").json()["unit"] == 1000000
    r = _post(client, h, bs_wb(), mode="commit", unit=1000)                      # 시트는 "원"
    assert r.status_code == 422 and any("단위" in e for e in _body(r)["errors"])
    r = _post(client, h, bs_wb(unit="(단위 : 억원)"), mode="commit")
    assert r.status_code == 422 and any("억원" in e for e in _body(r)["errors"])
    assert _count(tid, FsStatement) == 0


def test_decimal_amount_rejected_float_noise_accepted(client: TestClient) -> None:
    h, _ = _tenant(client)
    r = _post(client, h, horizontal_bs(stale_parent=False, cash_2025=120.5), mode="commit", unit=1)
    assert r.status_code == 422 and any("소수" in e for e in _body(r)["errors"])
    ok = _post(client, h, horizontal_bs(stale_parent=False, cash_2025=120.0000001), mode="commit", unit=1)
    assert ok.status_code == 200, ok.text


def test_orphan_value_warning(client: TestClient) -> None:
    h, _ = _tenant(client)
    b = _post(client, h, horizontal_bs(stale_parent=False, orphan=True), mode="preview", unit=1).json()
    assert any("계정명 없이" in w for w in b["warnings"])


def test_existing_statement_conflict_409(client: TestClient) -> None:
    h, tid = _tenant(client)
    assert _post(client, h, bs_wb(), mode="commit").status_code == 200
    n = _count(tid, FsAmount)
    r = _post(client, h, bs_wb(), mode="commit")
    assert r.status_code == 409
    b = _body(r)
    assert [c["fiscal_year"] for c in b["conflicts"]] == [2025] and b["conflicts"][0]["status"] == "final"
    assert _count(tid, FsAmount) == n


def test_existing_master_requires_mapping_then_reuses_accounts(client: TestClient) -> None:
    h, tid = _tenant(client)
    assert _post(client, h, bs_wb(), mode="commit").status_code == 200
    accounts = _count(tid, FsAccount)

    pre = _post(client, h, bs_wb(), mode="preview", fiscal_years="2024").json()
    assert pre["mapping_required"] and not pre["can_commit"] and not pre["master_empty"]
    sugg = pre["suggested_mapping"]
    assert len(sugg) == accounts and "new" not in sugg.values()                # 경로가 같아 전부 기존 계정
    assert pre["statements"][0]["ok"]                                          # 제안 대응으로 모의 검증

    assert _post(client, h, bs_wb(), mode="commit", fiscal_years="2024").status_code == 409
    ok = _post(client, h, bs_wb(), mode="commit", fiscal_years="2024", mapping=sugg)
    assert ok.status_code == 200, ok.text
    assert _count(tid, FsAccount) == accounts and _count(tid, FsStatement) == 2
    assert ok.json()["statements"][0]["status"] == "final"                     # 이번 업로드의 최신 연도

    bad = dict(sugg)
    bad.pop(next(iter(bad)))
    r = _post(client, h, bs_wb(), mode="preview", fiscal_years="2024", mapping=bad).json()
    assert any("mapping 에 없는 행" in e for e in r["errors"])


def test_include_prior_and_fiscal_year_selection(client: TestClient) -> None:
    h, _ = _tenant(client)
    b = _post(client, h, bs_wb(), mode="preview", include_prior="true").json()
    assert b["fiscal_years"] == [2025, 2024]
    assert [(s["fiscal_year"], s["finalize_candidate"]) for s in b["statements"]] == [(2025, True), (2024, False)]
    r = _post(client, h, bs_wb(), mode="preview", fiscal_years="2019").json()
    assert any("시트에 없는 회계연도" in e for e in r["errors"])


def test_sheet_selection_and_unsupported(client: TestClient) -> None:
    h, _ = _tenant(client)
    wb = _wb(("BS공시", "재 무 상 태 표", BS_ROWS), ("PL공시", "손 익 계 산 서", PL_ROWS))
    ce = wb.create_sheet("CE공시")
    ce["A1"] = "자 본 변 동 표"
    b = _post(client, h, wb, mode="preview").json()
    assert {s["sheet"] for s in b["sheets"]} == {"BS공시", "PL공시"} and "sheet 를 지정" in b["errors"][0]
    assert _post(client, h, wb, mode="commit").status_code == 422
    assert _post(client, h, wb, mode="preview", sheet="PL공시").json()["statement_type"] == "PL"
    r = _post(client, h, wb, mode="commit", sheet="CE공시")
    assert r.status_code == 422 and "자본변동표" in _body(r)["errors"][0]
    r = _post(client, h, wb, mode="preview", sheet="BS공시", statement_type="PL").json()
    assert any("다릅니다" in e for e in r["errors"])
    assert _post(client, h, wb, mode="preview", sheet="없음").status_code == 422


def test_rejects_non_xlsx_and_non_manager(client: TestClient) -> None:
    h, _ = _tenant(client)
    r = client.post("/api/fs/upload", headers=h, data={"mode": "preview"},
                    files={"file": ("a.csv", b"x,y", "text/csv")})
    assert r.status_code == 400
    r = client.post("/api/fs/upload", headers=h, data={"mode": "preview"},
                    files={"file": ("a.xlsx", b"not a zip", XLSX)})
    assert r.status_code == 400
    assert _post(client, h, bs_wb(), mode="save").status_code == 422
    viewer, _ = _tenant(client, roles=("external_auditor",))
    assert _post(client, viewer, bs_wb(), mode="preview").status_code == 403


def test_tenant_isolation_of_upload(client: TestClient) -> None:
    h1, t1 = _tenant(client)
    h2, t2 = _tenant(client)
    assert _post(client, h1, bs_wb(), mode="commit").status_code == 200
    b = _post(client, h2, bs_wb(), mode="preview").json()
    assert b["master_empty"] and not b["conflicts"]                             # 다른 테넌트 마스터는 안 보인다
    assert _count(t2, FsAccount) == 0 and _count(t1, FsAccount) == 13


# ── 8-B2 정산표 결합 ──────────────────────────────────────────
# 합성 공시 BS(BS_ROWS) 의 잎 공시 행과 금액이 맞는 합성 정산표. 매핑 열은 D.

def worksheet_bs(*, years=(2024, 2025), over: dict | None = None, remap: dict | None = None) -> Workbook:
    over, remap = over or {}, remap or {}
    rows = [  # (라벨, 굵게, 매핑, 2024, 2025)
        ("자산", True, None, None, None),
        ("Ⅰ.유동자산", True, None, 250, 300),
        ("현금", False, "현금및현금성자산", 100, 100),
        ("받을어음", False, "매출채권", 60, 80),
        ("외상매출금", False, "매출채권", 100, 130),
        ("대손충당금", False, "매출채권", -10, -10),
        ("재고자산", False, None, 0, 0),                        # 대응 없음·0 → 건너뜀
        ("Ⅱ.비유동자산", True, None, 450, 500),
        ("건물", False, "유형자산", 300, 300),
        ("감가상각누계액", False, "유형자산", -50, -60),
        ("비품", False, "유형자산", 220, 290),
        ("감가상각누계액", False, "유형자산", -20, -30),        # 같은 공시 행 아래 반복 이름
        ("자산총계", True, None, 700, 800),
        ("부채", True, None, None, None),
        ("Ⅰ.유동부채", True, None, 100, 120),
        ("외상매입금", False, "매입채무", 100, 120),
        ("부채총계", True, None, 100, 120),
        ("자본", True, None, None, None),
        ("Ⅰ.자본금", True, None, 500, 500),
        ("보통주자본금", False, "I. 자본금", 500, 500),          # 접두 번호가 붙은 매핑 값
        ("Ⅱ.이익잉여금", True, None, 100, 180),
        ("미처분이익잉여금", False, "이익잉여금", 100, 180),
        ("자본총계", True, None, 600, 680),
        ("부채및자본총계", True, None, 700, 800),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "BS정산표"
    ws.append(["과목", *years])
    for label, bold, mapto, a, b in rows:
        vals = {2024: a, 2025: b, 2026: b}
        if label in over:
            vals.update(over[label])
        ws.append([label, *[vals[y] for y in years], remap.get(label, mapto)])
        ws.cell(ws.max_row, 1).font = Font(bold=bold)
    return wb


def _attach(client, h, wb, **form):
    buf = BytesIO()
    wb.save(buf)
    return client.post("/api/fs/upload/attach", headers=h, data={k: str(v) for k, v in form.items()},
                       files={"file": ("ws.xlsx", buf.getvalue(), XLSX)})


def _disclosure_2y(client, h) -> dict[int, str]:
    r = _post(client, h, bs_wb(), mode="commit", include_prior="true")
    assert r.status_code == 200, r.text
    return {s["fiscal_year"]: s["statement_id"] for s in r.json()["statements"]}


def _reopen(client, h, sid) -> None:
    assert client.post(f"/api/fs/statements/{sid}/reopen", headers=h, json={"reason": "결합"}).status_code == 200


def test_attach_worksheet_makes_disclosure_lines_subtotals(client: TestClient) -> None:
    h, tid = _tenant(client)
    sids = _disclosure_2y(client, h)
    before = _count(tid, FsAccount)

    assert _attach(client, h, worksheet_bs(), mode="commit", unit=1).status_code == 409   # 2025 final
    pre = _attach(client, h, worksheet_bs(), mode="preview", unit=1).json()
    assert pre["errors"] == [] and pre["bridge_column"] == "D" and pre["fiscal_years"] == [2025, 2024]
    assert [s["ok"] for s in pre["statements"]] == [True, True]
    assert any("재고자산" in w for w in pre["warnings"])
    assert _count(tid, FsAccount) == before                                               # preview 저장 안 함

    _reopen(client, h, sids[2025])
    com = _attach(client, h, worksheet_bs(), mode="commit", unit=1)
    assert com.status_code == 200, com.text
    body = com.json()
    assert {s["fiscal_year"]: s["status"] for s in body["statements"]} == {2025: "final", 2024: "draft"}
    names = {r["name"] for r in body["rows"] if not r["skip_reason"]}
    assert {"감가상각누계액_건물", "감가상각누계액_비품", "보통주자본금"} <= names
    assert _count(tid, FsAccount) == before + 11

    d = client.get(f"/api/fs/statements/{sids[2024]}", headers=h).json()
    ar = next(n for n in d["tree"][0]["children"][0]["children"] if n["name"] == "매출채권")
    assert ar["is_subtotal"] and {c["name"] for c in ar["children"]} == {"받을어음", "외상매출금", "대손충당금"}
    bad_debt = next(c for c in ar["children"] if c["name"] == "대손충당금")
    assert bad_debt["raw_meta"]["attach"] and bad_debt["raw_meta"]["source_path"] == ["자산총계", "유동자산"]

    # 다시 붙이면 기존 COA 계정을 쓴다(계정 수 불변)
    _reopen(client, h, sids[2025])
    again = _attach(client, h, worksheet_bs(), mode="commit", unit=1)
    assert again.status_code == 200, again.text
    assert all(r["existing_account_id"] for r in again.json()["rows"] if not r["skip_reason"])
    assert _count(tid, FsAccount) == before + 11


def test_attach_mismatch_keeps_draft(client: TestClient) -> None:
    h, _ = _tenant(client)
    sids = _disclosure_2y(client, h)
    _reopen(client, h, sids[2025])
    r = _attach(client, h, worksheet_bs(over={"외상매출금": {2025: 131}}), mode="commit", unit=1)
    assert r.status_code == 200, r.text
    st = {s["fiscal_year"]: s for s in r.json()["statements"]}
    assert not st[2025]["ok"] and st[2025]["status"] == "draft"
    assert {(e["rule"], e["account_name"]) for e in st[2025]["errors"]} == {("suspense_unresolved", "임시계정(원본 차이)")}
    assert [(x["parent_name"], D(x["amount"])) for x in st[2025]["suspense"]] == [("매출채권", D(-1))]
    assert st[2024]["ok"]


def test_attach_unmatched_bridge_needs_category_map(client: TestClient) -> None:
    h, _ = _tenant(client)
    sids = _disclosure_2y(client, h)
    _reopen(client, h, sids[2025])
    wb = worksheet_bs(remap={"외상매입금": "매입채무및기타채무"})
    r = _attach(client, h, wb, mode="commit", unit=1)
    assert r.status_code == 422 and _body(r)["unmatched"] == ["매입채무및기타채무"]
    ok = _attach(client, h, wb, mode="commit", unit=1,
                 category_map=json.dumps({"매입채무및기타채무": "매입채무"}))
    assert ok.status_code == 200, ok.text


def test_attach_guards(client: TestClient) -> None:
    h, _ = _tenant(client)
    r = _attach(client, h, worksheet_bs(), mode="commit", unit=1)                  # 공시 재무제표 없음
    assert r.status_code == 422 and any("공시양식을 먼저" in e for e in _body(r)["errors"])

    sids = _disclosure_2y(client, h)
    _reopen(client, h, sids[2025])
    for form, word in (({}, "단위"), ({"unit": 1000}, "단위")):
        r = _attach(client, h, worksheet_bs(), mode="commit", **form)
        assert r.status_code == 422 and any(word in e for e in _body(r)["errors"])
    r = _attach(client, h, worksheet_bs(remap={"현금": "유동자산"}), mode="commit", unit=1)
    assert r.status_code == 422 and any("소계라" in e for e in _body(r)["errors"])
    # 2024 공시 재무제표에도 쓰이는 공시 행 — 정산표에 2024 열이 없으면 막는다(소계가 되면 2024 가 깨진다).
    # 연도 열이 하나면 가로 연도형이 아니므로 2025·2026 으로 만든다(2026 은 공시 재무제표가 없어 건너뜀)
    r = _attach(client, h, worksheet_bs(years=(2025, 2026)), mode="commit", unit=1)
    assert r.status_code == 422 and any("[2024]" in e for e in _body(r)["errors"])

    viewer, _ = _tenant(client, roles=("external_auditor",))
    assert _attach(client, viewer, worksheet_bs(), mode="preview", unit=1).status_code == 403


def test_suggestion_skips_account_whose_subtotal_flag_differs(client: TestClient) -> None:
    """같은 경로라도 소계 여부가 다르면 제안하지 않는다 — 결합으로 소계가 된 공시 행에 잎을 대응시키면 하위 0개 소계가 된다."""
    h, tid = _tenant(client)
    assert _post(client, h, horizontal_bs(stale_parent=False), mode="commit", unit=1).status_code == 200
    db = TestingSessionLocal()
    tok = set_active_tenant(tid)
    try:
        cash = db.scalars(select(FsAccount).where(FsAccount.name == "현금")).one()
        svc_fs.create_account(db, statement_type="BS", name="보통예금", section="asset", parent_id=cash.id)
        cash.is_subtotal = True            # 정산표 결합 뒤의 공시 행처럼 — 하위가 있는 소계
        db.commit()
        cash_id = str(cash.id)
    finally:
        reset_active_tenant(tok)
        db.close()
    pre = _post(client, h, horizontal_bs(stale_parent=False), mode="preview", unit=1, basis="consolidated").json()
    row = next(r for r in pre["rows"] if r["label"] == "현금")
    assert pre["suggested_mapping"][str(row["row_no"])] == "new"
    assert cash_id not in pre["suggested_mapping"].values()
    assert pre["suggested_mapping"][str(next(r for r in pre["rows"] if r["label"] == "토지")["row_no"])] != "new"
