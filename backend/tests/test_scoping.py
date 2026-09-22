"""스코핑 코어 검증 (6-1, ADR-0034).

**핵심 네 가지**(6-1 §5 — 1·3·7·10):
- **원천 재현** — 2022 수치를 넣으면 원천 캐시값과 한 원도 다르지 않다(픽스처로만 쓰고 시드하지 않는다)
- **원천 판정 재현** — 원천 질적 평가값으로 계산한 결론이 원천 결론과 같다. 평균 2.0 인 항목이 유의로 남는다
- **필드 단위 배지** — 한 필드를 고치면 그 필드만 떨어지고, 같은 값 재저장은 배지를 떼지 않는다
- **확정 스냅샷** — 확정 후 정책을 바꿔도 확정 당시 결론은 그대로다

원천 파일은 `backend/seeds/4__내부회계관리제도_Scoping.xlsx` 를 그대로 읽는다(`data_only=True` 캐시값).
검증 대상 동작은 HTTP 로 부른다(13.9-35 교훈). 전역 템플릿은 적재 함수로 한 번만 넣는다.
"""
import re
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.security import hash_password
from app.core.tenant_context import DEFAULT_TENANT_ID, reset_active_tenant, set_active_tenant
from app.models.scoping import (
    QUAL_COMPARISON_GE,
    QUAL_COMPARISON_GT,
    RATING_ALIASES,
    ScopingTemplateAccount,
    ScopingTemplateText,
)
from app.models.tenant import UserTenantAccess
from app.models.user import User
from app.models.user_mgmt import UserRole
from app.services import scoping_calc as calc
from seeds.seed_scoping_template import SOURCE, load_template
from tests.conftest import TestingSessionLocal

PW = "pw123456"
SHEETS = {"BS": "2.1 유의한 계정과목(BS)", "PL": "2.2 유의한 계정과목(PL)",
          "NOTE": "2.3 유의한 주석", "CF": "2.4 유의한 현금흐름 관련항목 "}


# ── 준비 ───────────────────────────────────────────────────

@pytest.fixture(scope="module", autouse=True)
def template(app):
    db = TestingSessionLocal()
    try:
        tpl, _, created = load_template(db)
        if created:
            db.commit()
        return tpl.id
    finally:
        db.close()


def _account(email: str, roles: tuple[str, ...] = ()) -> None:
    db = TestingSessionLocal()
    tok = set_active_tenant(DEFAULT_TENANT_ID)
    try:
        u = db.query(User).filter(User.email == email).first()
        if u is None:
            u = User(email=email, hashed_password=hash_password(PW), display_name=email.split("@")[0],
                     role="user", is_active=True)
            db.add(u)
            db.commit()
        if db.query(UserTenantAccess).filter(UserTenantAccess.user_id == u.id,
                                             UserTenantAccess.tenant_id == DEFAULT_TENANT_ID).first() is None:
            db.add(UserTenantAccess(user_id=u.id, tenant_id=DEFAULT_TENANT_ID, role="user"))
            db.commit()
        for r in roles:
            if db.query(UserRole).filter(UserRole.user_id == u.id, UserRole.role_name == r,
                                         UserRole.is_deleted == False).first() is None:  # noqa: E712
                db.add(UserRole(user_id=u.id, role_name=r))
        db.commit()
    finally:
        reset_active_tenant(tok)
        db.close()


def _headers(client: TestClient, email: str) -> dict:
    r = client.post("/api/auth/login", data={"username": email, "password": PW})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}


@pytest.fixture()
def mgr(client: TestClient) -> dict:
    _account("scope-mgr@acme.example", ("icfr_manager",))
    return _headers(client, "scope-mgr@acme.example")


def _create(client: TestClient, h: dict, year: int) -> dict:
    r = client.post("/api/scoping", headers=h, json={"fiscal_year": year})
    assert r.status_code == 201, r.text
    return r.json()


def _set_policy(client: TestClient, h: dict, key: str, value: str) -> None:
    r = client.put("/api/org/policies", headers=h, json={"policy_key": key, "policy_value": value})
    assert r.status_code == 200, r.text


def _acc(detail: dict, name: str) -> dict:
    return next(a for a in detail["accounts"] if a["name"] == name)


# ── 1. 원천 재현 ───────────────────────────────────────────

def test_reproduces_source_2022_materiality(client: TestClient, mgr: dict) -> None:
    """**2022 수치를 넣으면 원천 캐시값과 같다.** 픽스처로만 쓰고 시드하지 않는다."""
    src = load_workbook(SOURCE, data_only=True)["1. 중요성 기준"]
    assert (src["F54"].value, src["G58"].value, src["G62"].value) == (240942200, 240942000, 168659400)

    s = _create(client, mgr, 2092)
    r = client.patch(f"/api/scoping/{s['id']}/benchmarks/adjusted_pbt", headers=mgr,
                     json={"base_amount": src["F38"].value})
    assert r.status_code == 200, r.text
    d = r.json()
    pbt = next(b for b in d["benchmarks"] if b["kind"] == "adjusted_pbt")
    assert pbt["amount"] == src["F54"].value
    assert d["overall_materiality"] == src["G58"].value
    assert d["smt"] == src["G62"].value

    # 매출액 벤치마크도 원천과 같다 (G54 = ROUND(11141354865 × 0.005, 0))
    r = client.patch(f"/api/scoping/{s['id']}/benchmarks/revenue", headers=mgr,
                     json={"base_amount": src["F37"].value})
    rev = next(b for b in r.json()["benchmarks"] if b["kind"] == "revenue")
    assert rev["amount"] == src["G54"].value


# ── 2. 반올림 경계 ─────────────────────────────────────────

@pytest.mark.parametrize("value,digits,expected", [
    ("2.5", 0, "3"), ("-2.5", 0, "-3"), ("0.5", 0, "1"), ("-0.5", 0, "-1"), ("3.5", 0, "4"),
    ("2.4", 0, "2"), ("1500", -3, "2000"), ("-1500", -3, "-2000"), ("2500", -3, "3000"), ("499", -3, "0"),
])
def test_excel_round_half_away_from_zero(value: str, digits: int, expected: str) -> None:
    """Excel ROUND 는 0.5 를 0 에서 멀어지는 쪽으로 올린다. Python round() 는 은행가 반올림이라 다르다."""
    assert calc.excel_round(Decimal(value), digits) == Decimal(expected)


def test_python_round_would_differ() -> None:
    """이 테스트가 있는 이유 — round() 로 구현하면 한 원이 달라진다."""
    assert round(2.5) == 2 and calc.excel_round(Decimal("2.5"), 0) == 3


# ── 3·4. 원천 판정 재현 + 정책 전환 ────────────────────────

_REF = re.compile(r"(?<![A-Za-z])\$?[A-Z]{1,3}\$?\d+")


def _source_rows():
    """원천 계정 시트의 (종류, 계정명, 평가값, 양적(캐시), 결론 셀, 결론(캐시))."""
    wbf, wbv = load_workbook(SOURCE, data_only=False), load_workbook(SOURCE, data_only=True)
    for stype, sn in SHEETS.items():
        ws, wv = wbf[sn], wbv[sn]
        cols = {" ".join(str(ws.cell(17, c).value).split()): c
                for c in range(1, ws.max_column + 1) if ws.cell(17, c).value}
        q0 = next(c for n, c in cols.items() if n.startswith("질적기준"))
        concl = next(c for n, c in cols.items() if n.startswith("유의성"))
        quant = next((c for n, c in cols.items() if n.startswith("양적") and "유의" in n), None)
        amt = next(c for c in range(1, ws.max_column + 1) if re.match(r"FY\d{4}", str(wv.cell(17, c).value or "")))
        for r in range(18, ws.max_row + 1):
            name = ws.cell(r, 2).value
            if not name:
                continue
            a = ws.cell(r, amt).value
            if isinstance(a, str) and a.startswith("=") and "!" not in a and _REF.search(a[1:]):
                continue
            if ws.cell(r, concl).value is None:
                continue
            ratings = {f"q{i + 1}": RATING_ALIASES[str(ws.cell(r, q0 + i).value).strip().lower()]
                       for i in range(10) if ws.cell(r, q0 + i).value}
            qv = calc.NA if stype in ("NOTE", "CF") else wv.cell(r, quant).value
            yield stype, str(name).strip(), ratings, qv, ws.cell(r, concl).value, wv.cell(r, concl).value


def test_source_conclusions_reproduce_with_ge() -> None:
    """원천 질적 평가값으로 계산한 결론이 원천 결론과 같다(수식 행 전부). 평균 2.0 은 유의로 남는다."""
    matched = mismatched = 0
    eq2 = []
    for stype, name, ratings, qv, raw, cached in _source_rows():
        if not (isinstance(raw, str) and raw.startswith("=")):
            continue   # 직접 입력(수동 판정)은 따로 본다
        qual = calc.qualitative(ratings, Decimal(2), QUAL_COMPARISON_GE)
        ours = calc.conclusion(qv, qual)
        matched += ours == cached
        mismatched += ours != cached
        if calc.qualitative_average(ratings) == Decimal(2) and qv not in ("Y",):
            eq2.append(name)
    assert mismatched == 0
    assert matched == 185
    assert len(eq2) == 16   # "초과"로 바꾸면 비유의로 떨어지는 수식 행


def test_gt_policy_flips_exactly_the_16() -> None:
    """정책을 "초과"로 바꾸면 그 16건이 비유의로 바뀐다 — 설정이 동작한다."""
    flipped = []
    for stype, name, ratings, qv, raw, _ in _source_rows():
        if not (isinstance(raw, str) and raw.startswith("=")):
            continue
        ge = calc.conclusion(qv, calc.qualitative(ratings, Decimal(2), QUAL_COMPARISON_GE))
        gt = calc.conclusion(qv, calc.qualitative(ratings, Decimal(2), QUAL_COMPARISON_GT))
        if ge == "Y" and gt == "N":
            flipped.append(name)
    assert len(flipped) == 16
    assert {"대손충당금(매출채권)", "미수수익", "리스부채", "이연법인세자산"} <= set(flipped)


def test_policy_switch_applies_through_api(client: TestClient, mgr: dict) -> None:
    """API 에서도 정책이 판정을 바꾼다(하드코딩 아님). 대손충당금은 평균 2.0 이다."""
    s = _create(client, mgr, 2093)
    try:
        d = client.get(f"/api/scoping/{s['id']}", headers=mgr).json()
        assert _acc(d, "대손충당금(매출채권)")["qual"] == "Y"
        _set_policy(client, mgr, "scoping_qual_comparison", "gt")
        d = client.get(f"/api/scoping/{s['id']}", headers=mgr).json()
        assert _acc(d, "대손충당금(매출채권)")["qual"] == "N"
    finally:
        _set_policy(client, mgr, "scoping_qual_comparison", "ge")


def test_policy_values_are_validated(client: TestClient, mgr: dict) -> None:
    for key, value in (("scoping_qual_comparison", "over"), ("scoping_qual_threshold", "5")):
        r = client.put("/api/org/policies", headers=mgr, json={"policy_key": key, "policy_value": value})
        assert r.status_code == 422, r.text


# ── 5. 미입력 → 미평가 ─────────────────────────────────────

def test_incomplete_qual_is_unevaluated_not_zero() -> None:
    """10요소 중 하나라도 없으면 질적 미평가. 원천처럼 빈칸을 0 으로 합산하지 않는다."""
    ratings = {f"q{i}": "H" for i in range(1, 10)}   # q10 없음
    assert calc.qualitative(ratings, Decimal(2), QUAL_COMPARISON_GE) is None
    assert calc.conclusion("N", None) is None            # 양적 N + 질적 미평가 = 미평가
    assert calc.conclusion("Y", None) == "Y"
    assert calc.conclusion(calc.NA, None) is None        # 주석·현금흐름 — 질적만으로
    assert calc.quantitative("NOTE", 10**12, 1) == calc.NA    # 금액이 있어도 해당 없음
    assert calc.quantitative("BS", None, 1) is None           # 미입력 = 미평가 (해당 없음과 다르다)


# ── 6. 수동 판정 ───────────────────────────────────────────

def test_manual_requires_reason_and_keeps_both(client: TestClient, mgr: dict) -> None:
    s = _create(client, mgr, 2094)
    d = client.get(f"/api/scoping/{s['id']}", headers=mgr).json()
    a = _acc(d, "대손충당금(매출채권)")
    r = client.patch(f"/api/scoping/{s['id']}/accounts/{a['id']}", headers=mgr, json={"manual_conclusion": "N"})
    assert r.status_code == 422, r.text
    r = client.patch(f"/api/scoping/{s['id']}/accounts/{a['id']}", headers=mgr,
                     json={"manual_conclusion": "N", "manual_reason": "검토 결과 중요하지 않음"})
    a = _acc(r.json(), "대손충당금(매출채권)")
    assert (a["computed"], a["manual_conclusion"], a["final"]) == ("Y", "N", "N")


def test_source_manual_judgements_are_seeded() -> None:
    """원천에서 계산 결론과 다르게 직접 입력된 6건만 수동 판정이다."""
    db = TestingSessionLocal()
    try:
        manual = {a.name for a in db.query(ScopingTemplateAccount).filter(
            ScopingTemplateAccount.manual_conclusion.isnot(None))}
    finally:
        db.close()
    assert manual == {"보통주자본금", "주식발행초과금", "기타자본잉여금", "자기주식", "미처분이익잉여금", "14. 무형자산"}


# ── 7. 필드 단위 배지 ──────────────────────────────────────

def test_badges_are_field_level(client: TestClient, mgr: dict) -> None:
    """**생성 직후 템플릿 값 전부 배지. 한 필드 수정 → 그 필드만 해제. 같은 값 재저장 → 해제 안 됨.**"""
    s = _create(client, mgr, 2095)
    d = client.get(f"/api/scoping/{s['id']}", headers=mgr).json()
    a = _acc(d, "대손충당금(매출채권)")
    assert all(v == "template" for v in a["badges"].values()) and len(a["badges"]) == 11  # 10요소 + 판단 근거(수동 판정 없음)
    assert all(t["badge"] == "template" for t in d["texts"])
    before = d["badge_count"]

    # 같은 값 재저장 — 배지 그대로
    same = a["ratings"]["q1"]
    d = client.patch(f"/api/scoping/{s['id']}/accounts/{a['id']}", headers=mgr,
                     json={"ratings": {"q1": same}}).json()
    assert _acc(d, "대손충당금(매출채권)")["badges"]["ratings.q1"] == "template"
    assert d["badge_count"] == before

    # 한 요소만 바꿈 — 그 요소만 떨어짐
    new = "L" if same != "L" else "M"
    d = client.patch(f"/api/scoping/{s['id']}/accounts/{a['id']}", headers=mgr,
                     json={"ratings": {"q1": new}}).json()
    badges = _acc(d, "대손충당금(매출채권)")["badges"]
    assert badges["ratings.q1"] == "edited"
    assert all(v == "template" for k, v in badges.items() if k != "ratings.q1")
    assert d["badge_count"] == before - 1


# ── 8·9·10. 확정 ───────────────────────────────────────────

def test_confirm_records_badges_and_reason_then_locks(client: TestClient, mgr: dict) -> None:
    s = _create(client, mgr, 2096)
    base = f"/api/scoping/{s['id']}"
    assert client.post(f"{base}/transition", headers=mgr, json={"to_status": "review"}).status_code == 200
    assert client.post(f"{base}/transition", headers=mgr,
                       json={"to_status": "confirmed"}).status_code == 422   # 사유 없음
    d = client.post(f"{base}/transition", headers=mgr,
                    json={"to_status": "confirmed", "reason": "감사인 검토 완료"}).json()
    assert d["status"] == "confirmed"
    assert d["confirm_reason"] == "감사인 검토 완료"
    assert d["confirm_badge_count"] == d["badge_count"] > 0   # 경고로 남는다, 막지는 않는다

    # 확정 후 수정 거부
    a = d["accounts"][0]
    assert client.patch(f"{base}/accounts/{a['id']}", headers=mgr,
                        json={"qual_basis": "x"}).status_code == 409
    # 재오픈 — 사유 필수, 이력
    assert client.post(f"{base}/transition", headers=mgr, json={"to_status": "draft"}).status_code == 422
    d = client.post(f"{base}/transition", headers=mgr,
                    json={"to_status": "draft", "reason": "금액 재검토"}).json()
    assert d["status"] == "draft" and d["confirmed_snapshot"] is None
    assert [(h["from_status"], h["to_status"]) for h in d["history"]] == [
        ("draft", "review"), ("review", "confirmed"), ("confirmed", "draft")]
    assert d["history"][-1]["reason"] == "금액 재검토"


def test_snapshot_survives_policy_change(client: TestClient, mgr: dict) -> None:
    """**확정 후 정책 기준을 바꿔도 확정 스냅샷의 결론은 그대로다.**"""
    s = _create(client, mgr, 2097)
    base = f"/api/scoping/{s['id']}"
    client.post(f"{base}/transition", headers=mgr, json={"to_status": "review"})
    d = client.post(f"{base}/transition", headers=mgr,
                    json={"to_status": "confirmed", "reason": "확정"}).json()
    aid = _acc(d, "대손충당금(매출채권)")["id"]
    assert d["confirmed_snapshot"]["accounts"][aid]["final"] == "Y"
    try:
        _set_policy(client, mgr, "scoping_qual_comparison", "gt")
        d = client.get(base, headers=mgr).json()
        a = _acc(d, "대손충당금(매출채권)")
        # 현재 정책으로 산출하면 바뀐다 — 질적 N, 금액을 넣지 않아 양적 미평가 → 결론 미평가
        assert a["qual"] == "N" and a["final"] is None
        assert a["snapshot_final"] == "Y"    # 확정 당시 판단은 그대로
        assert d["confirmed_snapshot"]["policy"]["comparison"] == "ge"
    finally:
        _set_policy(client, mgr, "scoping_qual_comparison", "ge")


def test_summary_counts_snapshot_when_confirmed(client: TestClient, mgr: dict) -> None:
    """대시보드 집계도 확정 후에는 스냅샷으로 센다 — 계정 표와 같은 기준. 미평가는 N 과 따로 센다."""
    s = _create(client, mgr, 2100)   # 테스트 중 가장 큰 연도 — 요약은 최근 연도를 본다
    base = f"/api/scoping/{s['id']}"
    before = client.get("/api/scoping/summary", headers=mgr).json()
    assert before["fiscal_year"] == 2100 and before["status"] == "draft"
    client.post(f"{base}/transition", headers=mgr, json={"to_status": "review"})
    client.post(f"{base}/transition", headers=mgr, json={"to_status": "confirmed", "reason": "확정"})
    try:
        _set_policy(client, mgr, "scoping_qual_comparison", "gt")
        after = client.get("/api/scoping/summary", headers=mgr).json()
        assert after["by_statement"] == before["by_statement"]
        assert sum(c["unevaluated"] for c in after["by_statement"].values()) > 0
    finally:
        _set_policy(client, mgr, "scoping_qual_comparison", "ge")


# ── 11. 가이드 범위 ────────────────────────────────────────

def test_out_of_range_warns_but_saves(client: TestClient, mgr: dict) -> None:
    s = _create(client, mgr, 2098)
    base = f"/api/scoping/{s['id']}"
    d = client.patch(f"{base}/benchmarks/adjusted_pbt", headers=mgr, json={"rate": "0.2"}).json()
    assert next(b for b in d["benchmarks"] if b["kind"] == "adjusted_pbt")["out_of_range"] is True
    assert any("조정세전순이익" in w for w in d["warnings"])
    d = client.patch(base, headers=mgr, json={"smt_rate": "0.9"}).json()
    assert d["smt_rate"] == "0.9" and d["smt_out_of_range"] is True
    # 매출액은 원천 범위가 오기 의심이라 경고하지 않는다
    d = client.patch(f"{base}/benchmarks/revenue", headers=mgr, json={"rate": "0.5"}).json()
    assert next(b for b in d["benchmarks"] if b["kind"] == "revenue")["out_of_range"] is False


# ── 12. 템플릿 내용 ────────────────────────────────────────

def test_template_has_no_amounts_emails_or_old_headers() -> None:
    db = TestingSessionLocal()
    try:
        texts = [t.body + (t.title or "") for t in db.query(ScopingTemplateText)]
        accs = db.query(ScopingTemplateAccount).all()
    finally:
        db.close()
    blob = "\n".join(texts + [" ".join(filter(None, (a.name, a.group_label, a.qual_basis, a.manual_reason)))
                              for a in accs])
    assert "@" not in blob
    assert "FY2019" not in blob and "4818843993" not in blob and "4,818,843,993" not in blob
    assert not re.search(r"사이냅|synap|deloitte", blob, re.I)
    assert not hasattr(ScopingTemplateAccount, "current_amount")   # 템플릿에는 금액 컬럼 자체가 없다
    assert "이상일 경우" in blob and "초과할 경우" not in blob     # "초과" → "이상" 반영


# ── 13. 권한 ───────────────────────────────────────────────

def test_external_auditor_and_plain_user_cannot_write(client: TestClient, mgr: dict) -> None:
    s = _create(client, mgr, 2099)
    _account("scope-ext@acme.example", ("external_auditor",))
    _account("scope-plain@acme.example")
    for email in ("scope-ext@acme.example", "scope-plain@acme.example"):
        h = _headers(client, email)
        assert client.get(f"/api/scoping/{s['id']}", headers=h).status_code == 200
        assert client.post("/api/scoping", headers=h, json={"fiscal_year": 2091}).status_code == 403
        assert client.patch(f"/api/scoping/{s['id']}", headers=h, json={"rationale": "x"}).status_code == 403
        assert client.post(f"/api/scoping/{s['id']}/transition", headers=h,
                           json={"to_status": "review"}).status_code == 403


def test_duplicate_year_and_recreate(client: TestClient, mgr: dict) -> None:
    """같은 회계연도는 1건. 409 문구는 원인을 말한다."""
    _create(client, mgr, 2090)
    r = client.post("/api/scoping", headers=mgr, json={"fiscal_year": 2090})
    assert r.status_code == 409 and "2090" in r.json()["detail"]
