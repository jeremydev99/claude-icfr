"""스코핑 엑셀 갱신 업로드 (13.9-118) — 미리보기·맞추기·반영·잠금.

원천 양식(`seeds/4__…Scoping.xlsx`)을 고쳐 올린다. openpyxl 로 저장하면 수식 셀의 캐시값이 사라지므로
테스트에 쓰는 행은 결론을 **값**으로 써 둔다(실제 엑셀 저장본은 캐시값이 있다). 회사 실 파일은 저장소에 넣지 않는다 —
`SCOPING_SAMPLE_FILE` 이 있으면 로컬에서만 실 파일 파싱을 확인한다.
"""
import json
import os
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services import scoping_excel as sx
from seeds.seed_scoping_template import SOURCE
from tests.test_scoping import (  # noqa: F401
    _acc,
    _create,
    _reopen,
    _submit_and_approve,
    lead,
    mgr,
    template,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BS, NOTE = "2.1 유의한 계정과목(BS)", "2.3 유의한 주석"


def _row(ws, name: str) -> int:
    return next(r for r in range(18, ws.max_row + 1) if str(ws.cell(r, 2).value or "").strip() == name)


def _book(edit) -> bytes:
    wb = load_workbook(SOURCE)
    edit(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _edit(wb) -> None:
    ws = wb[BS]
    r = _row(ws, "단기금융상품")                 # 10요소 빈칸 → 전부 Low + 근거
    for c in range(6, 16):
        ws.cell(r, c).value = "Low"
    ws.cell(r, 17).value = "테스트 근거"
    ws.cell(r, 18).value = "N"
    r = _row(ws, "미수금")                       # 이름이 달라도 괄호를 빼면 같다
    ws.cell(r, 2).value = "미수금(테스트)"
    ws.cell(r, 17).value = "괄호 근거"
    ws.cell(r, 18).value = "N"
    r = _row(ws, "선급금")                       # 아무 규칙에도 안 맞는 이름 → 사람이 고른다
    ws.cell(r, 2).value = "완전히다른이름"
    ws.cell(r, 17).value = "직접 연결 근거"
    ws.cell(r, 18).value = "N"
    wn = wb[NOTE]                                # 주석은 양적 해당 없음 — 결론 = 질적. 다르게 쓰면 수동 판정
    r = _row(wn, "1. 당사의 개요")
    wn.cell(r, 18).value = "회사 개요 수동 근거"
    wn.cell(r, 19).value = "Y"


def _post(client, h, sid, content, mode, links=None):
    data = {"mode": mode}
    if links:
        data["links"] = json.dumps(links)
    return client.post(f"/api/scoping/{sid}/upload-excel", headers=h, data=data,
                       files={"file": ("scoping.xlsx", content, XLSX)})


def test_preview_matches_and_commit_applies_as_edited(client: TestClient, mgr: dict) -> None:  # noqa: F811
    d = _create(client, mgr, 2041)
    content = _book(_edit)
    r = _post(client, mgr, d["id"], content, "preview")
    assert r.status_code == 200, r.text
    p = r.json()
    by = {c["account_name"]: c for c in p["changes"]}
    assert {f["field"] for f in by["단기금융상품"]["fields"]} == {"ratings", "qual_basis"}
    assert by["단기금융상품"]["fields"][0]["after"] == "LLLLLLLLLL"
    assert by["미수금"]["how"] == "괄호·번호 제외" and by["미수금"]["excel_name"] == "미수금(테스트)"
    assert by["1. 회사의 개요"]["how"].startswith("이름")          # '당사' → '회사' 로 맞춘다
    assert {f["field"] for f in by["1. 회사의 개요"]["fields"]} >= {"manual"}
    unmatched = {x["name"]: x for x in p["excel_only"]}
    assert "완전히다른이름" in unmatched and "선급금" in {a["name"] for a in p["account_only"]}
    assert p["summary_text"].startswith("미리보기")
    # 미리보기는 아무것도 바꾸지 않는다
    got = client.get(f"/api/scoping/{d['id']}", headers=mgr).json()
    assert _acc(got, "단기금융상품")["ratings"] == {}

    # 사람이 고른 연결과 함께 반영
    sel = next(a for a in p["account_only"] if a["name"] == "선급금")
    r = _post(client, mgr, d["id"], content, "commit", {unmatched["완전히다른이름"]["key"]: sel["id"]})
    assert r.status_code == 200, r.text
    assert any(c["how"] == "직접 연결" and c["account_name"] == "선급금" for c in r.json()["changes"])
    got = client.get(f"/api/scoping/{d['id']}", headers=mgr).json()
    a = _acc(got, "단기금융상품")
    assert a["ratings"] == {f"q{i}": "L" for i in range(1, 11)} and a["qual_basis"] == "테스트 근거"
    assert "template" not in a["badges"].values()
    assert _acc(got, "선급금")["qual_basis"] == "직접 연결 근거"
    n = _acc(got, "1. 회사의 개요")
    assert n["manual_conclusion"] == "Y" and n["manual_reason"] == "회사 개요 수동 근거"
    assert n["badges"]["qual_basis"] == "edited"          # 템플릿 값을 고친 칸 = 수정됨(화면 수정과 같다)
    # 같은 파일을 다시 올리면 바뀌는 계정이 없다(연결도 같게)
    again = _post(client, mgr, d["id"], content, "preview",
                  {unmatched["완전히다른이름"]["key"]: sel["id"]}).json()
    assert again["changed_count"] == 0


def test_not_applicable_is_skipped_and_amounts_not_applied(client: TestClient, mgr: dict) -> None:  # noqa: F811
    d = _create(client, mgr, 2042)
    a = _acc(d, "단기금융상품")
    assert client.patch(f"/api/scoping/{d['id']}/accounts/{a['id']}", headers=mgr,
                        json={"not_applicable": True, "na_reason": "거래 없음"}).status_code == 200
    p = _post(client, mgr, d["id"], _book(_edit), "preview").json()
    assert "단기금융상품" not in {c["account_name"] for c in p["changes"]}
    assert any(s["account_name"] == "단기금융상품" for s in p["skipped"])
    got = client.get(f"/api/scoping/{d['id']}", headers=mgr).json()
    assert _acc(got, "단기금융상품")["current_amount"] is None     # 금액은 반영하지 않는다


def test_commit_refused_unless_draft(client: TestClient, mgr: dict, lead: dict) -> None:  # noqa: F811
    d = _create(client, lead, 2043)
    base = f"/api/scoping/{d['id']}"
    assert client.patch(base, headers=lead, json={"rationale": "근거"}).status_code == 200
    _submit_and_approve(client, base, lead, mgr, "확정")
    content = _book(_edit)
    assert _post(client, lead, d["id"], content, "preview").status_code == 200
    assert _post(client, lead, d["id"], content, "commit").status_code == 409


def test_bad_rating_is_reported() -> None:
    def bad(wb):
        ws = wb[BS]
        ws.cell(_row(ws, "단기금융상품"), 6).value = "Huge"
        ws.cell(_row(ws, "단기금융상품"), 18).value = "N"
    with pytest.raises(sx.ExcelError, match="Huge"):
        sx.parse(_book(bad))


@pytest.mark.skipif(not os.environ.get("SCOPING_SAMPLE_FILE"), reason="로컬 전용 — 회사 실 파일")
def test_real_sample_parses() -> None:
    with open(os.environ["SCOPING_SAMPLE_FILE"], "rb") as f:
        rows, _ = sx.parse(f.read())
    counts = {k: len(v) for k, v in rows.items()}
    assert all(counts[k] > 0 for k in ("BS", "PL", "NOTE", "CF")), counts
    assert not any("총계" in r.name for v in rows.values() for r in v)


# ── 직접 연결 기억·현금흐름 표기 (13.9-146) ─────────────────

def _edit_basis(text: str):
    def edit(wb) -> None:
        _edit(wb)
        ws = wb[BS]
        ws.cell(_row(ws, "완전히다른이름"), 17).value = text
    return edit


def test_direct_link_is_remembered_next_year_and_can_be_overridden(client: TestClient, mgr: dict) -> None:  # noqa: F811
    """반영한 직접 연결은 이름 기준으로 기억 → 다음 해 스코핑에서 사람 선택 없이 '기억한 연결'. 다시 고르면 덮어쓴다."""
    d = _create(client, mgr, 2071)
    content = _book(_edit)
    p = _post(client, mgr, d["id"], content, "preview").json()
    # 앞 테스트가 같은 행을 이미 기억시켰을 수 있다 — 엑셀 키는 미매칭·변경 어느 쪽에서든 찾는다
    key = next(x["key"] for x in p["excel_only"] + p["changes"]
               if x.get("name", x.get("excel_name")) == "완전히다른이름")
    sel = _acc(client.get(f"/api/scoping/{d['id']}", headers=mgr).json(), "선급금")
    assert _post(client, mgr, d["id"], content, "commit", {key: sel["id"]}).status_code == 200

    nxt = _create(client, mgr, 2072)                                    # 다음 해 — 계정 행이 새로 만들어진다
    r = _post(client, mgr, nxt["id"], _book(_edit_basis("내년 근거")), "preview").json()
    by = {c["account_name"]: c for c in r["changes"]}
    assert by["선급금"]["how"] == "기억한 연결" and r["remembered"] >= 1
    assert "완전히다른이름" not in {x["name"] for x in r["excel_only"]}

    # 다른 계정으로 다시 고르면 덮어쓴다 — 그다음부터는 새 계정
    other = next(a for a in r["account_only"] if a["name"] != "선급금")
    assert _post(client, mgr, nxt["id"], _book(_edit_basis("내년 근거")), "commit", {key: other["id"]}).status_code == 200
    again = _post(client, mgr, nxt["id"], _book(_edit_basis("다시 근거")), "preview").json()
    hit = next(c for c in again["changes"] if c["excel_name"] == "완전히다른이름")
    assert hit["how"] == "기억한 연결" and hit["account_name"] == other["name"]


def test_cash_flow_name_variants() -> None:
    v = sx._cf_variants
    assert v("5. 법인세의 환급(납부)") & v("법인세의 납부")
    assert v("매출채권의 감소(증가)") == v("매출채권의 증가(감소)")
    assert v("토지의 취득") == v("토지의 증가") and v("자기주식의 처분") == v("자기주식의 감소")
    assert v("V. 기초의 현금") == v("기초의 현금및현금성자산")
    assert not (v("대여금의 감소") & v("차입금의 감소"))
